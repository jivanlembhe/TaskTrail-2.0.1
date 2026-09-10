#!/usr/bin/env python3
"""
TaskTrail — task planner, to-do list and status tracker (Python / PySide6 edition)

Data-compatible with the Electron TaskTrail: same JSON format, same file
(%APPDATA%\\TaskTrail\\flowboard_data.json on Windows), same backups.

Run:     python tasktrail.py
Build:   see README.md (build_win.bat / GitHub Actions)
"""
import calendar
import copy
import datetime as dt
import json
import os
import re
import shutil
import sys
import time
import uuid

from PySide6.QtCore import (Qt, QTimer, QSize, QRect, QRectF, QPoint, Signal, QDate, QEvent, QMimeData, QVariantAnimation,
                            QPropertyAnimation, QEasingCurve, QAbstractAnimation)
from PySide6.QtGui import (QAction, QColor, QFont, QIcon, QPainter, QPixmap, QBrush, QPen,
                           QFontDatabase, QCursor, QShortcut, QKeySequence, QDrag)
from PySide6.QtWidgets import (QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout, QGridLayout,
                               QLabel, QPushButton, QFrame, QStackedWidget, QListWidget, QListWidgetItem,
                               QAbstractItemView, QScrollArea, QLineEdit, QComboBox, QTextEdit, QDialog,
                               QDialogButtonBox, QCheckBox, QDateEdit, QColorDialog, QFileDialog, QMessageBox,
                               QSystemTrayIcon, QMenu, QSizePolicy, QToolButton, QDockWidget, QSpinBox,
                               QFontComboBox, QSlider, QInputDialog, QSplitter, QGraphicsOpacityEffect)

APP_NAME = "TaskTrail"
VERSION = "2.2.0"
MONTHS = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
MONTHS_LONG = ['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August',
               'September', 'October', 'November', 'December']
DEFAULT_COLS = [
    {"id": "todo", "label": "To Do", "color": "#6c8aff", "icon": "📋"},
    {"id": "today", "label": "Today's", "color": "#ffb347", "icon": "☀️"},
    {"id": "wip", "label": "WIP", "color": "#b48aff", "icon": "⚡"},
    {"id": "done", "label": "Done", "color": "#3ddbbf", "icon": "✅"},
]
BUILTIN = {"todo", "today", "wip", "done"}
LEGACY_YEAR = 2026
LABEL_PALETTE = ['#6c8aff', '#ffb347', '#b48aff', '#3ddbbf', '#ff6584', '#26de81',
                 '#fd9644', '#a55eea', '#2bcbba', '#eb3b5a', '#f7b731', '#778ca3']
PRI_LABEL = {"high": "🔴 High", "med": "🟡 Med", "low": "🟢 Low"}
PRI_COLOR = {"high": "#ff6584", "med": "#ffb347", "low": "#3ddbbf"}

THEMES = {
    "dark": dict(bg="#07080d", s1="#0e1018", s2="#13161f", s3="#191d29", s4="#1f2333", border="#252a3a",
                 border2="#2e3448", text="#eef0f8", text2="#9aa0bc", text3="#5c6382", accent="#6c8aff"),
    "light": dict(bg="#f0f2f9", s1="#ffffff", s2="#ffffff", s3="#f4f5fb", s4="#eaecf5", border="#d8dbee",
                  border2="#c4c8e0", text="#1a1d2e", text2="#4a5070", text3="#8890b0", accent="#4a6bef"),
}
PRESETS = [
    ("Default dark", "dark", {}),
    ("Default light", "light", {}),
    ("Midnight", "dark", dict(bg="#0a1020", surface="#121a30", text="#e8ecff", accent="#7c9cff")),
    ("Forest", "dark", dict(bg="#0b1410", surface="#12211a", text="#e6f2ec", accent="#3ddbbf")),
    ("Ember", "dark", dict(bg="#16100d", surface="#241a16", text="#f7ece6", accent="#ff8c5a")),
    ("Paper", "light", dict(bg="#f5f2ea", surface="#fffdf8", text="#25231f", accent="#b5622e")),
    ("Ocean", "dark", dict(bg="#061420", surface="#0d2233", text="#e3f1fb", accent="#35c2ff")),
    ("Rose", "dark", dict(bg="#1a0f14", surface="#26161d", text="#fbe9ef", accent="#ff6584")),
    ("Sunset", "dark", dict(bg="#1a1208", surface="#2a1f10", text="#fff1dd", accent="#ffb347")),
    ("Graphite", "dark", dict(bg="#141414", surface="#1e1e1e", text="#ececec", accent="#9aa0bc")),
    ("Violet", "dark", dict(bg="#120d1f", surface="#1b1430", text="#efe9ff", accent="#b48aff")),
    ("Lavender", "light", dict(bg="#f3f0fa", surface="#ffffff", text="#2a2340", accent="#7c5cd6")),
    ("Mint", "light", dict(bg="#eef7f3", surface="#ffffff", text="#16302a", accent="#1fa27a")),
    ("Sky", "light", dict(bg="#eaf3fb", surface="#ffffff", text="#14243a", accent="#2b7de9")),
    ("Slate", "light", dict(bg="#e9ecf1", surface="#f7f9fc", text="#1f2937", accent="#4f6b8f")),
    ("Peach", "light", dict(bg="#fdf1ea", surface="#fffaf7", text="#3a2a22", accent="#f26b4e")),
]


# ═══════════════════════════════════════════════════════════════════════════
#  PATHS & SETTINGS
# ═══════════════════════════════════════════════════════════════════════════
def user_data_dir():
    if sys.platform.startswith("win"):
        base = os.environ.get("APPDATA") or os.path.expanduser("~")
    elif sys.platform == "darwin":
        base = os.path.expanduser("~/Library/Application Support")
    else:
        base = os.environ.get("XDG_CONFIG_HOME") or os.path.expanduser("~/.config")
    d = os.path.join(base, APP_NAME)
    os.makedirs(d, exist_ok=True)
    return d


USER_DATA = user_data_dir()
DB_FILE = os.path.join(USER_DATA, "flowboard_data.json")
SETTINGS_FILE = os.path.join(USER_DATA, "settings.json")
DEFAULT_BACKUP_DIR = os.path.join(USER_DATA, "backups")


def read_settings():
    try:
        with open(SETTINGS_FILE, encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}


def write_settings(s):
    with open(SETTINGS_FILE, "w", encoding="utf-8") as f:
        json.dump(s, f, indent=2)


def get_backup_dir():
    d = read_settings().get("backupDir") or DEFAULT_BACKUP_DIR
    try:
        os.makedirs(d, exist_ok=True)
        return d
    except Exception:
        os.makedirs(DEFAULT_BACKUP_DIR, exist_ok=True)
        return DEFAULT_BACKUP_DIR


def detect_onedrive():
    for p in (os.environ.get("OneDriveCommercial"), os.environ.get("OneDriveConsumer"),
              os.environ.get("OneDrive"), os.path.join(os.path.expanduser("~"), "OneDrive")):
        if p and os.path.isdir(p):
            return p
    return None


def rgba(hex_color, a):
    """Qt stylesheets don't accept #RRGGBBAA — build rgba() instead."""
    h = hex_color.lstrip('#')
    if len(h) != 6:
        return hex_color
    r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    return f"rgba({r},{g},{b},{a})"


def uid():
    return uuid.uuid4().hex[:12]


def now_iso():
    return dt.datetime.now().isoformat(timespec="seconds")


def fmt_date(s):
    try:
        return dt.date.fromisoformat(s).strftime("%b %d").replace(" 0", " ")
    except Exception:
        return s or ""


def fmt_dt(s):
    try:
        return dt.datetime.fromisoformat(s.replace("Z", "")).strftime("%b %d, %H:%M")
    except Exception:
        return ""


def is_overdue(d, done):
    if not d or done:
        return False
    try:
        return dt.date.fromisoformat(d) < dt.date.today()
    except Exception:
        return False


REPEATS = [("", "No repeat"), ("daily", "Every day"), ("weekdays", "Every weekday"), ("weekly", "Every week"), ("monthly", "Every month"), ("yearly", "Every year")]
REPEAT_SHORT = {"daily": "daily", "weekdays": "weekdays", "weekly": "weekly", "monthly": "monthly", "yearly": "yearly"}


def repeat_step(d, repeat):
    """The occurrence after date `d` for a repeat rule."""
    if repeat == "daily": return d + dt.timedelta(1)
    if repeat == "weekdays":
        d += dt.timedelta(1)
        while d.weekday() > 4: d += dt.timedelta(1)
        return d
    if repeat == "weekly": return d + dt.timedelta(7)
    if repeat == "monthly":
        y, m = (d.year + 1, 1) if d.month == 12 else (d.year, d.month + 1)
        return dt.date(y, m, min(d.day, calendar.monthrange(y, m)[1]))
    if repeat == "yearly": return dt.date(d.year + 1, d.month, min(d.day, calendar.monthrange(d.year + 1, d.month)[1]))
    return None


def next_due(iso, repeat, after=None):
    """Next occurrence strictly after `after` (default: the due date itself). None if the rule is empty/invalid."""
    try: d = dt.date.fromisoformat(iso)
    except (TypeError, ValueError): return None
    after = after or d
    while d <= after:
        d = repeat_step(d, repeat)
        if d is None: return None
    return d


def occurrences(card, start, end):
    """Projected future dates of a recurring card inside [start, end], after its current due date (which is shown as the real card)."""
    if not card.get("repeat") or not card.get("dueDate"): return []
    try: d = dt.date.fromisoformat(card["dueDate"]); until = dt.date.fromisoformat(card["repeatUntil"]) if card.get("repeatUntil") else None
    except ValueError: return []
    out = []
    while True:
        d = repeat_step(d, card["repeat"])
        if d is None or d > end or (until and d > until) or len(out) > 62: break
        if d >= start: out.append(d)
    return out


DOW = ['mon', 'tue', 'wed', 'thu', 'fri', 'sat', 'sun']     # date.weekday(): Mon = 0


def parse_date_token(tok, year, month):
    """@today · @tomorrow · @+3 · @fri · @15 (day of the shown month) · @2026-10-01 → ISO date or None."""
    t = tok.lower(); today = dt.date.today()
    if t in ("today", "tod"): return today.isoformat()
    if t in ("tomorrow", "tmr", "tom"): return (today + dt.timedelta(1)).isoformat()
    if re.fullmatch(r"\+\d+", t): return (today + dt.timedelta(int(t[1:]))).isoformat()
    di = next((i for i, d in enumerate(DOW) if t.startswith(d)), -1)
    if di >= 0: return (today + dt.timedelta((di - today.weekday()) % 7)).isoformat()
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", t):
        try: return dt.date.fromisoformat(t).isoformat()
        except ValueError: return None
    if re.fullmatch(r"\d{1,2}", t) and 1 <= int(t) <= calendar.monthrange(year, month + 1)[1]:
        return dt.date(year, month + 1, int(t)).isoformat()
    return None


def parse_quick(store, text, create=True):
    """'Renew SSL !high #ops @fri' → dict(title, priority, dueDate, labels). #label creates the label if needed
    (create=False: don't create, list the would-be-new names under 'new' instead — used by the live preview)."""
    out = dict(title="", priority="med", dueDate="", labels=[]); keep = []
    for w in text.split():
        sig, val = w[0], w[1:]
        if sig not in "!#@*" or not val: keep.append(w); continue
        v = val.lower()
        if sig == "*":
            rp = {"daily": "daily", "day": "daily", "weekdays": "weekdays", "weekday": "weekdays", "weekly": "weekly", "week": "weekly", "monthly": "monthly", "month": "monthly", "yearly": "yearly", "year": "yearly"}.get(v)
            if rp: out["repeat"] = rp
            else: keep.append(w)
        elif sig == "!":
            pri = {"high": "high", "h": "high", "hi": "high", "urgent": "high", "low": "low", "l": "low", "med": "med", "m": "med", "medium": "med"}.get(v)
            if pri: out["priority"] = pri
            else: keep.append(w)
        elif sig == "@":
            d = parse_date_token(val, store.year, store.month)
            if d: out["dueDate"] = d
            else: keep.append(w)
        else:
            l = next((x for x in store.labels if x["name"].lower() == v), None)
            if not l and not create:
                if val not in out.setdefault("new", []): out["new"].append(val)
                continue
            if not l:
                l = {"id": uid(), "name": val, "color": LABEL_PALETTE[len(store.labels) % len(LABEL_PALETTE)]}; store.labels.append(l)
            if l["id"] not in out["labels"]: out["labels"].append(l["id"])
    out["title"] = " ".join(keep).strip()
    if out.get("repeat") and not out["dueDate"]: out["dueDate"] = dt.date.today().isoformat()      # a repeat needs a date to count from
    return out


# ═══════════════════════════════════════════════════════════════════════════
#  STORE — same JSON schema as the Electron app
# ═══════════════════════════════════════════════════════════════════════════
class Store:
    def __init__(self):
        self.db = {}
        self.year = dt.date.today().year
        self.month = dt.date.today().month - 1
        self.load()
        if not os.path.exists(DB_FILE):
            self.save()                 # make sure the file exists from the first run so backups work

    # ---- load / save ----
    def load(self):
        raw = None
        if os.path.exists(DB_FILE):
            try:
                with open(DB_FILE, encoding="utf-8") as f:
                    raw = json.load(f)
            except Exception:
                raw = None
        self.db = self.normalize(raw)
        self.ensure_year(self.year); self.merged_count = 0
        if not self.db["meta"].get("checklistMerged"):        # one-time: checklist → board (backup of the old file first)
            if raw is not None: self.backup_now()
            self.merged_count = self.merge_checklist(); self.save()

    @staticmethod
    def normalize(raw):
        d = raw if isinstance(raw, dict) else {}
        if not isinstance(d.get("years"), dict):
            d["years"] = {}
            if d.get("kanban") or d.get("checklist"):
                d["years"][str(LEGACY_YEAR)] = {"kanban": d.get("kanban") or {}, "checklist": d.get("checklist") or {}}
        d.pop("kanban", None)
        d.pop("checklist", None)
        if not isinstance(d.get("columns"), list) or not d["columns"]:
            d["columns"] = copy.deepcopy(DEFAULT_COLS)
        for dc in DEFAULT_COLS:
            if not any(c["id"] == dc["id"] for c in d["columns"]):
                d["columns"].append(dict(dc))
        d.setdefault("labels", [])
        d.setdefault("migrations", [])
        for m in d["migrations"]:
            m.setdefault("fromYear", LEGACY_YEAR)
            m.setdefault("toYear", LEGACY_YEAR)
        if not isinstance(d.get("meta"), dict):
            d["meta"] = {}
        return d

    def ensure_year(self, y):
        Y = self.db["years"].setdefault(str(y), {"kanban": {}})
        Y.setdefault("kanban", {})
        for mi in range(12):
            k = Y["kanban"].setdefault(str(mi), {})
            for c in self.db["columns"]:
                k.setdefault(c["id"], [])
        return Y

    def merge_checklist(self):
        """Checklist groups become labels; their tasks become board cards (To Do, or Done if ticked). Runs once, also for restored old backups."""
        moved = 0
        for y, Y in list(self.db["years"].items()):
            for mi, groups in list((Y.get("checklist") or {}).items()):
                for g in groups or []:
                    if not g.get("tasks"): continue
                    l = next((x for x in self.labels if x["name"].lower() == g["name"].strip().lower()), None)
                    if not l:
                        l = {"id": uid(), "name": g["name"].strip(), "color": LABEL_PALETTE[len(self.labels) % len(LABEL_PALETTE)]}; self.labels.append(l)
                    kd = self.kanban(int(mi), int(y))
                    for t in g["tasks"]:
                        done = bool(t.get("done"))
                        card = dict(id=t.get("id") or uid(), title=t.get("title", ""), desc="", priority=t.get("priority", "med"), dueDate=t.get("dueDate", ""), done=done,
                                    subs=t.get("subs", []), labels=[l["id"]], comments=[], activity=[], createdAt=t.get("createdAt") or now_iso(), month=MONTHS[int(mi)], year=int(y))
                        if t.get("migratedFrom"): card["migratedFrom"] = t["migratedFrom"]
                        self.log(card, f'Merged from checklist group "{g["name"]}"'); kd.setdefault("done" if done else "todo", []).append(card); moved += 1
            Y.pop("checklist", None)
        self.db["meta"]["checklistMerged"] = now_iso(); self.db["meta"]["checklistMergedCount"] = moved
        return moved

    def complete_recurring(self, card):
        """Ticking a recurring card advances its due date instead of finishing it. Returns the new date, or None if the rule ended."""
        today = dt.date.today()
        try: due = dt.date.fromisoformat(card["dueDate"])
        except (TypeError, ValueError): return None
        nd = next_due(card["dueDate"], card["repeat"], after=max(due, today))
        if nd is None or (card.get("repeatUntil") and nd.isoformat() > card["repeatUntil"]):
            card["repeat"] = ""; return None
        card.setdefault("completions", []).append(today.isoformat()); del card["completions"][:-120]
        card["dueDate"] = nd.isoformat(); card["updatedAt"] = now_iso(); self.log(card, f"Completed ↻ next due {fmt_date(card['dueDate'])}")
        return card["dueDate"]

    def save(self):
        self.db["meta"]["updatedAt"] = int(time.time() * 1000)
        tmp = DB_FILE + ".tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(self.db, f)
        os.replace(tmp, DB_FILE)

    # ---- accessors ----
    @property
    def columns(self):
        return self.db["columns"]

    @property
    def labels(self):
        return self.db["labels"]

    def col(self, cid):
        return next((c for c in self.columns if c["id"] == cid), None)

    def kanban(self, mi=None, year=None):
        Y = self.ensure_year(year or self.year)
        return Y["kanban"][str(self.month if mi is None else mi)]

    def years(self):
        return sorted(int(y) for y in self.db["years"])

    def find_card(self, cid, col_id=None):
        kd = self.kanban()
        for c in self.columns:
            if col_id and c["id"] != col_id:
                continue
            for card in kd.get(c["id"], []):
                if card["id"] == cid:
                    return card, c["id"]
        return None, None

    @staticmethod
    def log(card, text):
        card.setdefault("activity", []).append({"at": now_iso(), "text": text})
        del card["activity"][:-60]

    def move_card(self, cid, from_col, to_col, before_id="TOP"):
        src = self.kanban()[from_col]
        idx = next((i for i, c in enumerate(src) if c["id"] == cid), -1)
        if idx < 0:
            return None
        card = src.pop(idx)
        dst = self.kanban().setdefault(to_col, [])
        if before_id == "TOP":
            at = 0
        elif before_id is None:
            at = len(dst)
        else:
            at = next((i for i, c in enumerate(dst) if c["id"] == before_id), len(dst))
        dst.insert(at, card)
        if to_col == "done" and from_col != "done" and card.get("repeat") and card.get("dueDate"):
            if self.complete_recurring(card):                       # stays in its column with the next date
                dst.pop(at); src.insert(idx, card); card["recurred"] = True; return card
        card.pop("recurred", None)
        if from_col != to_col:
            card["done"] = to_col == "done"
            if to_col == "done":
                card["completedAt"] = now_iso()
            self.log(card, f"Moved from {self.col(from_col)['label']} to {self.col(to_col)['label']}")
        return card

    # ---- rollover ----
    def incomplete_col_ids(self):
        return [c["id"] for c in self.columns if c["id"] != "done"]

    def incomplete_cards(self, mi, year):
        kd = self.kanban(mi, year)
        return [c for cid in self.incomplete_col_ids() for c in kd.get(cid, []) if not c.get("done")]

    @staticmethod
    def month_past(mi, year):
        last = dt.date(year + (1 if mi == 11 else 0), 1 if mi == 11 else mi + 2, 1) - dt.timedelta(days=1)
        return dt.date.today() > last

    def past_periods(self):
        today = dt.date.today()
        out = []
        for y in self.years():
            if y > today.year:
                continue
            last_m = 11 if y < today.year else today.month - 2
            for mi in range(0, last_m + 1):
                if self.month_past(mi, y):
                    out.append((y, mi))
        return out

    def migrate_month(self, mi, year, auto=False):
        ty, tm = (year + 1, 0) if mi >= 11 else (year, mi + 1)
        self.ensure_year(year); self.ensure_year(ty)
        src, tgt = self.kanban(mi, year), self.kanban(tm, ty)
        label = f"{MONTHS[mi]} {year}"
        moved_k = moved_c = 0
        for cid in self.incomplete_col_ids():
            arr = src.get(cid, [])
            for c in [c for c in arr if not c.get("done")]:
                m = copy.deepcopy(c)
                m.update(id=uid(), migratedFrom=label, migratedAt=now_iso(), originalId=c["id"])
                self.log(m, f"Rolled over from {label}")
                tgt.setdefault("todo" if cid == "today" else cid, []).insert(0, m)
                moved_k += 1
            src[cid] = [c for c in arr if c.get("done")]
        self.db["migrations"].append({"fromYear": year, "from": mi, "toYear": ty, "to": tm, "kanban": moved_k,
                                      "checklist": moved_c, "at": now_iso(), "auto": auto})
        self.save()
        return moved_k, moved_c

    def auto_migrate(self):
        total = 0
        for y, mi in self.past_periods():
            if not self.incomplete_cards(mi, y):
                continue
            if any(m.get("fromYear") == y and m["from"] == mi and m.get("auto") for m in self.db["migrations"]):
                continue
            k, c = self.migrate_month(mi, y, auto=True)
            total += k + c
        return total

    def pending_rollover(self):
        return any(self.incomplete_cards(mi, y) for y, mi in self.past_periods())

    # ---- backups ----
    def backup_now(self):
        if not os.path.exists(DB_FILE):
            return None
        d = get_backup_dir()
        dest = os.path.join(d, "flowboard_backup_" + dt.datetime.now().strftime("%Y-%m-%dT%H-%M-%S-%f")[:-3] + ".json")
        shutil.copyfile(DB_FILE, dest)
        files = sorted([f for f in os.listdir(d) if f.startswith("flowboard_backup_") and f.endswith(".json")],
                       key=lambda f: os.path.getmtime(os.path.join(d, f)), reverse=True)
        for f in files[30:]:
            try:
                os.remove(os.path.join(d, f))
            except Exception:
                pass
        return dest

    @staticmethod
    def backup_list():
        d = get_backup_dir()
        out = []
        for f in os.listdir(d):
            if f.endswith(".json"):
                p = os.path.join(d, f)
                out.append((f, os.path.getsize(p), os.path.getmtime(p)))
        return sorted(out, key=lambda x: x[2], reverse=True)

    def restore_file(self, path):
        with open(path, encoding="utf-8") as f:
            parsed = json.load(f)
        if not isinstance(parsed, dict) or not (parsed.get("kanban") or parsed.get("years")):
            raise ValueError("Not a TaskTrail / FlowBoard backup file.")
        data = open(path, "rb").read()          # read first: the safety backup must never clobber the source
        self.backup_now()
        with open(DB_FILE, "wb") as f:
            f.write(data)
        self.load()
        self.save()


# ═══════════════════════════════════════════════════════════════════════════
#  THEME / STYLESHEET
# ═══════════════════════════════════════════════════════════════════════════
def build_qss(p, font_family, font_size):
    return f"""
    QWidget {{ background:{p['bg']}; color:{p['text']}; font-family:"{font_family}"; font-size:{font_size}px; }}
    QMainWindow, QDialog {{ background:{p['bg']}; }}
    QFrame#sidebar {{ background:{p['s1']}; border-right:1px solid {p['border']}; }}
    QFrame#topbar {{ background:{p['s1']}; border-bottom:1px solid {p['border']}; }}
    QLabel#brand {{ font-size:{font_size+8}px; font-weight:800; color:{p['accent']}; }}
    QLabel#brandSub, QLabel#navSection {{ color:{p['text3']}; font-size:{font_size-4}px; letter-spacing:1px; }}
    QLabel#pageTitle {{ font-size:{font_size+5}px; font-weight:700; }}
    QLabel#pageMonth {{ color:{p['accent']}; background:{rgba(p['accent'], 0.14)}; border-radius:10px; font-weight:600; }}
    QPushButton#nav {{ text-align:left; padding:8px 18px; border:none; border-left:3px solid transparent; color:{p['text2']}; background:transparent; font-weight:500; }}
    QPushButton#nav:hover {{ color:{p['text']}; background:{p['s2']}; }}
    QPushButton#nav:checked {{ color:{p['accent']}; background:{rgba(p['accent'], 0.12)}; border-left:3px solid {p['accent']}; }}
    QPushButton#month {{ padding:4px 2px; border:1px solid {p['border']}; border-radius:6px; color:{p['text3']}; background:transparent; font-size:{font_size-3}px; font-weight:700; }}
    QPushButton#month:hover {{ border-color:{p['accent']}; color:{p['accent']}; }}
    QPushButton#month:checked {{ background:{p['accent']}; border-color:{p['accent']}; color:#ffffff; }}
    QPushButton {{ padding:6px 14px; border-radius:14px; border:1.5px solid {p['border']}; background:transparent; color:{p['text2']}; font-weight:700; }}
    QPushButton:hover {{ color:{p['text']}; border-color:{p['text2']}; }}
    QPushButton#primary {{ background:{p['accent']}; border-color:{p['accent']}; color:#ffffff; }}
    QPushButton#primary:hover {{ background:{p['accent']}dd; }}
    QPushButton#danger {{ color:#ff6584; border-color:#ff6584; }}
    QPushButton#ghost {{ border:none; color:{p['text3']}; padding:2px 6px; }}
    QPushButton#ghost:hover {{ color:{p['text']}; background:{p['s3']}; border-radius:6px; }}
    QPushButton#addk {{ border:1.5px dashed {p['border']}; border-radius:12px; color:{p['text3']}; font-weight:500; text-align:left; padding:8px 14px; }}
    QPushButton#addk:hover {{ border-color:{p['accent']}; color:{p['accent']}; }}
    QLineEdit, QTextEdit, QComboBox, QDateEdit, QSpinBox, QFontComboBox {{ background:{p['s3']}; border:1.5px solid {p['border']}; border-radius:8px; padding:6px 10px; color:{p['text']}; selection-background-color:{p['accent']}; }}
    QLineEdit:focus, QTextEdit:focus, QComboBox:focus, QDateEdit:focus {{ border-color:{p['accent']}; }}
    QComboBox QAbstractItemView {{ background:{p['s2']}; color:{p['text']}; selection-background-color:{p['accent']}; border:1px solid {p['border']}; }}
    QFrame#panel, QFrame#kpi, QFrame#card, QFrame#colhead, QFrame#group {{ background:{p['s2']}; border:1px solid {p['border']}; border-radius:12px; }}
    QFrame#card:hover {{ border-color:{p['border2']}; }}
    QFrame#detail {{ background:{p['s1']}; border-left:1px solid {p['border']}; }}
    QListWidget {{ background:transparent; border:none; outline:0; }}
    QListWidget::item {{ background:transparent; border:none; padding:0; margin:0 0 6px 0; }}
    QListWidget::item:selected {{ background:transparent; }}
    QScrollArea {{ border:none; background:transparent; }}
    QScrollBar:vertical {{ width:10px; background:transparent; }} QScrollBar:horizontal {{ height:10px; background:transparent; }}
    QScrollBar::handle {{ background:{p['border2']}; border-radius:5px; min-height:24px; min-width:24px; }}
    QScrollBar::handle:hover {{ background:{p['text3']}; }}
    QScrollBar::add-line, QScrollBar::sub-line {{ height:0; width:0; }}
    QLabel, QCheckBox {{ background:transparent; }}
    QLabel#muted {{ color:{p['text3']}; }} QLabel#muted2 {{ color:{p['text2']}; }}
    QLabel#kpiVal {{ font-size:{font_size+16}px; font-weight:700; }}
    QLabel#kpiLabel {{ color:{p['text2']}; font-size:{font_size-3}px; letter-spacing:1px; }}
    QLabel#sectitle {{ font-weight:700; font-size:{font_size+1}px; }}
    QLabel#badge {{ border-radius:9px; padding:1px 7px; font-size:{font_size-4}px; font-weight:700; }}
    QCheckBox {{ spacing:8px; }} QCheckBox::indicator {{ width:16px; height:16px; border-radius:4px; border:2px solid {p['border2']}; background:transparent; }}
    QCheckBox::indicator:checked {{ background:#3ddbbf; border-color:#3ddbbf; }}
    QMenu {{ background:{p['s2']}; border:1px solid {p['border']}; }} QMenu::item:selected {{ background:{p['accent']}; color:#fff; }}
    QToolTip {{ background:{p['s2']}; color:{p['text']}; border:1px solid {p['border']}; }}
    QDockWidget::title {{ background:{p['s1']}; padding:6px; }}
    QFrame#calday {{ background:{p['s2']}; border:1px solid {p['border']}; border-radius:10px; }}
    QFrame#calday[out="true"] {{ background:transparent; }}
    QFrame#calday[today="true"] {{ border:2px solid {p['accent']}; }}
    QFrame#calday[over="true"] {{ border:2px dashed {p['accent']}; background:{rgba(p['accent'], 0.10)}; }}
    """


# ═══════════════════════════════════════════════════════════════════════════
#  SMALL WIDGETS
# ═══════════════════════════════════════════════════════════════════════════
def badge(text, color, bg=None):
    l = QLabel(text)
    l.setObjectName("badge")
    l.setStyleSheet(f"QLabel#badge{{color:{color};background:{bg or rgba(color, 0.15)};}}")
    l.setMargin(2)
    return l


def hline():
    f = QFrame(); f.setFrameShape(QFrame.HLine); f.setStyleSheet("color: palette(mid);")
    return f


class BarChart(QWidget):
    """Tasks per month (current year). Click a bar to open that month; the shown month is highlighted."""
    clicked = Signal(int)

    def __init__(self, get_values, accent):
        super().__init__(); self.get_values = get_values; self.accent = accent; self.current = -1; self.progress = 1.0
        self.setMinimumHeight(110); self.setCursor(QCursor(Qt.PointingHandCursor))

    def paintEvent(self, e):
        vals = self.get_values(); mx = max(1, max(vals))
        p = QPainter(self); p.setRenderHint(QPainter.Antialiasing)
        w = self.width(); h = self.height() - 20; bw = (w - 11 * 6) / 12
        for i, v in enumerate(vals):
            bh = max(4, int(v / mx * (h - 8) * self.progress)); x = int(i * (bw + 6)); c = QColor(self.accent)
            if i != self.current: c.setAlphaF(0.42)
            p.setPen(Qt.NoPen); p.setBrush(c); p.drawRoundedRect(x, h - bh, int(bw), bh, 3, 3)
            p.setPen(QColor(self.accent if i == self.current else "#5c6382")); p.drawText(QRect(x, h + 2, int(bw), 16), Qt.AlignCenter, MONTHS[i][0])
        p.end()

    def mousePressEvent(self, e):
        i = int(e.position().x() / (self.width() / 12))
        if 0 <= i < 12: self.clicked.emit(i)


class Ring(QWidget):
    """Completion ring (0–100). `value` is animated by the dashboard."""
    def __init__(self, color):
        super().__init__(); self.color = color; self.value = 0; self.setFixedSize(56, 56)

    def paintEvent(self, e):
        p = QPainter(self); p.setRenderHint(QPainter.Antialiasing); r = QRectF(5, 5, 46, 46)
        pen = QPen(QColor(self.color)); pen.setWidth(6); pen.setCapStyle(Qt.RoundCap); track = QColor(self.color); track.setAlphaF(0.18)
        p.setPen(QPen(track, 6)); p.drawEllipse(r); p.setPen(pen); p.drawArc(r, 90 * 16, -int(360 * 16 * self.value / 100))
        p.setPen(QColor(self.color)); f = p.font(); f.setBold(True); f.setPixelSize(12); p.setFont(f); p.drawText(r, Qt.AlignCenter, f"{int(self.value)}%"); p.end()


class PriMix(QWidget):
    """Open tasks by priority as one stacked bar."""
    def __init__(self):
        super().__init__(); self.counts = (0, 0, 0); self.setFixedHeight(10)

    def paintEvent(self, e):
        tot = sum(self.counts); p = QPainter(self); p.setRenderHint(QPainter.Antialiasing); p.setPen(Qt.NoPen); x = 0.0
        if not tot: p.setBrush(QColor("#5c6382")); p.setOpacity(0.25); p.drawRoundedRect(QRectF(0, 0, self.width(), 10), 5, 5); p.end(); return
        for n, c in zip(self.counts, (PRI_COLOR["high"], PRI_COLOR["med"], PRI_COLOR["low"])):
            if n: w = self.width() * n / tot; p.setBrush(QColor(c)); p.drawRoundedRect(QRectF(x, 0, w, 10), 5, 5); x += w
        p.end()


class CardWidget(QFrame):
    """A task card on the board."""
    clicked = Signal(str, str)      # cid, col
    edit = Signal(str, str)
    delete = Signal(str, str)
    toggle = Signal(str, str)

    def __init__(self, store, card, col):
        super().__init__(); self.setObjectName("card"); self.card = card; self.col_id = col["id"]
        self.setCursor(QCursor(Qt.PointingHandCursor))
        self.setStyleSheet(f"QFrame#card{{border-left:3px solid {col['color']};}}")
        v = QVBoxLayout(self); v.setContentsMargins(12, 10, 10, 10); v.setSpacing(6)
        labels = [l for l in store.labels if l["id"] in card.get("labels", [])]
        if labels:
            lr = QHBoxLayout(); lr.setSpacing(4)
            for l in labels:
                lr.addWidget(badge(l["name"], "#ffffff", l["color"]))
            lr.addStretch(); v.addLayout(lr)
        top = QHBoxLayout(); top.setSpacing(8)
        done = card.get("done") or self.col_id == "done"
        cb = QCheckBox(); cb.setChecked(done); cb.setToolTip("Complete / reopen")
        cb.clicked.connect(lambda: self.toggle.emit(card["id"], self.col_id))
        title = QLabel(card["title"]); title.setWordWrap(True); title.setStyleSheet("font-weight:500;" + ("text-decoration:line-through;color:#5c6382;" if done else ""))
        title.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Preferred)
        eb = QPushButton("✎"); eb.setObjectName("ghost"); eb.setToolTip("Edit"); eb.clicked.connect(lambda: self.edit.emit(card["id"], self.col_id))
        db_ = QPushButton("✕"); db_.setObjectName("ghost"); db_.setToolTip("Delete"); db_.clicked.connect(lambda: self.delete.emit(card["id"], self.col_id))
        for w in (cb, title, eb, db_):
            top.addWidget(w)
        v.addLayout(top)
        meta = QHBoxLayout(); meta.setSpacing(5)
        meta.addWidget(badge(PRI_LABEL[card.get("priority", "med")], PRI_COLOR[card.get("priority", "med")]))
        if card.get("dueDate"):
            od = is_overdue(card["dueDate"], done)
            meta.addWidget(badge(("⚠ " if od else "📅 ") + fmt_date(card["dueDate"]), "#ff6584" if od else "#8890b0"))
        meta.addStretch(); v.addLayout(meta)
        if card.get("repeat") or card.get("migratedFrom"):      # second row so narrow columns don't clip the badges
            m2 = QHBoxLayout(); m2.setSpacing(5)
            if card.get("repeat"): m2.addWidget(badge("↻ " + REPEAT_SHORT.get(card["repeat"], card["repeat"]), "#b48aff"))
            if card.get("migratedFrom"): m2.addWidget(badge("↪ " + card["migratedFrom"], "#8890b0"))
            m2.addStretch(); v.addLayout(m2)
        if card.get("desc"):
            d = QLabel(card["desc"][:160] + ("…" if len(card["desc"]) > 160 else "")); d.setObjectName("muted2"); d.setWordWrap(True); v.addWidget(d)
        subs = card.get("subs", [])
        if subs:
            sd = sum(1 for s in subs if s.get("done"))
            v.addWidget(QLabel(f"☑ {sd}/{len(subs)} sub-tasks", objectName="muted"))
        if card.get("comments"):
            v.addWidget(QLabel(f"💬 {len(card['comments'])}", objectName="muted"))

    def mouseReleaseEvent(self, e):
        if e.button() == Qt.LeftButton:
            self.clicked.emit(self.card["id"], self.col_id)
        super().mouseReleaseEvent(e)


class CardList(QListWidget):
    """Column list with drag & drop between/within columns. The drop updates the model
    and asks the board to re-render, so item widgets are always rebuilt."""
    dropped = Signal(str, str, str, object)   # cid, from_col, to_col, before_id(None=end)

    def __init__(self, col_id):
        super().__init__(); self.col_id = col_id
        self.setDragDropMode(QAbstractItemView.DragDrop); self.setDefaultDropAction(Qt.MoveAction)
        self.setAcceptDrops(True); self.setDragEnabled(True); self.setSelectionMode(QAbstractItemView.SingleSelection)
        self.setVerticalScrollMode(QAbstractItemView.ScrollPerPixel); self.setSpacing(0)
        self.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        self.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        self.setUniformItemSizes(False)

    def resizeEvent(self, e):
        super().resizeEvent(e)
        self.fit_items()

    def fit_items(self):
        """Size every card to the column width so wrapped text and borders are never clipped."""
        w = self.viewport().width() - 2
        if w < 60:
            return
        for i in range(self.count()):
            it = self.item(i)
            wd = self.itemWidget(it)
            if wd is None:
                continue
            wd.setFixedWidth(w)
            h = wd.heightForWidth(w) if wd.layout() and wd.layout().hasHeightForWidth() else wd.sizeHint().height()
            it.setSizeHint(QSize(w, max(h, wd.sizeHint().height()) + 6))

    def dropEvent(self, e):
        src = e.source()
        if not isinstance(src, CardList) or not src.currentItem():
            e.ignore(); return
        cid = src.currentItem().data(Qt.UserRole)
        pos = e.position().toPoint() if hasattr(e, "position") else e.pos()
        item = self.itemAt(pos)
        before = None
        if item is not None:
            r = self.visualItemRect(item)
            if pos.y() < r.center().y():
                before = item.data(Qt.UserRole)
            else:
                nxt = self.item(self.row(item) + 1)
                before = nxt.data(Qt.UserRole) if nxt else None
        e.setDropAction(Qt.IgnoreAction); e.accept()
        self.dropped.emit(cid, src.col_id, self.col_id, before)


class ListRow(QFrame):
    """One task in the list view — the same card as on the board, one line."""
    clicked = Signal(str, str); toggle = Signal(str, str)

    def __init__(self, store, card, col):
        super().__init__(); self.setObjectName("card"); self.card = card; self.col_id = col["id"]; self.setCursor(QCursor(Qt.PointingHandCursor))
        self.setStyleSheet(f"QFrame#card{{border-left:3px solid {col['color']};}}"); h = QHBoxLayout(self); h.setContentsMargins(10, 6, 10, 6); h.setSpacing(8)
        done = card.get("done") or col["id"] == "done"; cb = QCheckBox(); cb.setChecked(done); cb.setToolTip("Complete / reopen"); cb.clicked.connect(lambda: self.toggle.emit(card["id"], col["id"])); h.addWidget(cb)
        t = QLabel(card["title"]); t.setWordWrap(True); t.setStyleSheet("font-weight:500;" + ("text-decoration:line-through;color:#5c6382;" if done else "")); h.addWidget(t, 1)
        subs = card.get("subs", [])
        if subs: h.addWidget(badge(f"☑ {sum(1 for x in subs if x.get('done'))}/{len(subs)}", "#8890b0"))
        for l in (x for x in store.labels if x["id"] in card.get("labels", [])): h.addWidget(badge(l["name"], "#ffffff", l["color"]))
        if card.get("repeat"): h.addWidget(badge("↻ " + REPEAT_SHORT.get(card["repeat"], card["repeat"]), "#b48aff"))
        if card.get("dueDate"): od = is_overdue(card["dueDate"], done); h.addWidget(badge(("⚠ " if od else "📅 ") + fmt_date(card["dueDate"]), "#ff6584" if od else "#8890b0"))
        h.addWidget(badge(PRI_LABEL[card.get("priority", "med")], PRI_COLOR[card.get("priority", "med")]))
        h.addWidget(badge(f"{col.get('icon', '')} {col['label']}", "#ffffff", col["color"]))

    def mouseReleaseEvent(self, e):
        if e.button() == Qt.LeftButton: self.clicked.emit(self.card["id"], self.col_id)


class CalChip(QLabel):
    """A task on the calendar. Board cards (drag_id set) can be dragged onto a day or the Unscheduled tray."""
    def __init__(self, text, color, drag_id=None, done=False, overdue=False, dashed=False, on_click=None, ghost=False):
        super().__init__(text); self.drag_id = drag_id; self.on_click = on_click; self._press = None
        self.setToolTip(text + (" — drag to reschedule" if drag_id else " — future occurrence; complete the task to move it here" if ghost else " — click to open")); self.setCursor(QCursor(Qt.PointingHandCursor))
        self.setSizePolicy(QSizePolicy.Ignored, QSizePolicy.Fixed)      # clip long titles instead of widening the day
        self.setStyleSheet(f"QLabel{{border-left:3px {'dashed' if dashed else 'solid'} {color};border-radius:5px;padding:2px 6px;background:{rgba(color, 0.05 if ghost else 0.14)};"
                           + ("color:#8890b0;" if ghost else "text-decoration:line-through;color:#5c6382;" if done else "color:#ff6584;font-weight:600;" if overdue else "") + "}")

    def mousePressEvent(self, e):
        if e.button() == Qt.LeftButton: self._press = e.position().toPoint()

    def mouseMoveEvent(self, e):
        if self._press is None or not self.drag_id or (e.position().toPoint() - self._press).manhattanLength() < QApplication.startDragDistance(): return
        self._press = None; md = QMimeData(); md.setText(self.drag_id); d = QDrag(self); d.setMimeData(md); d.setPixmap(self.grab()); d.exec(Qt.MoveAction)

    def mouseReleaseEvent(self, e):
        if self._press is not None and self.on_click: self.on_click()
        self._press = None


class CalDay(QFrame):
    """A calendar day (or, with iso='', the Unscheduled tray) that accepts dropped CalChips."""
    dropped = Signal(str, str)      # cid, iso ('' clears the due date)
    added = Signal(str)             # iso — double-clicked

    def __init__(self, iso):
        super().__init__(); self.iso = iso; self.setObjectName("calday"); self.setAcceptDrops(True)

    def _over(self, on):
        self.setProperty("over", on); self.style().unpolish(self); self.style().polish(self)

    def dragEnterEvent(self, e):
        if e.mimeData().hasText(): e.acceptProposedAction(); self._over(True)

    def dragLeaveEvent(self, e): self._over(False)

    def dropEvent(self, e):
        self._over(False); e.acceptProposedAction(); self.dropped.emit(e.mimeData().text(), self.iso)

    def mouseDoubleClickEvent(self, e):
        if self.iso: self.added.emit(self.iso)


# ═══════════════════════════════════════════════════════════════════════════
#  DIALOGS
# ═══════════════════════════════════════════════════════════════════════════
class TaskDialog(QDialog):
    def __init__(self, parent, store, card=None, col_id="todo", preset=None):
        super().__init__(parent); self.store = store; self.card = card
        self.setWindowTitle("Edit Task" if card else "New Task"); self.setMinimumWidth(540)
        v = QVBoxLayout(self); v.setSpacing(10)
        self.title = QLineEdit(card["title"] if card else ""); self.title.setPlaceholderText("What needs to be done?")
        self.desc = QTextEdit(card.get("desc", "") if card else ""); self.desc.setPlaceholderText("Details or notes…"); self.desc.setFixedHeight(70)
        self.col = QComboBox()
        for c in store.columns:
            self.col.addItem(f"{c.get('icon', '')} {c['label']}", c["id"])
        self.col.setCurrentIndex(max(0, self.col.findData(col_id)))
        self.pri = QComboBox()
        for k, lab in PRI_LABEL.items():
            self.pri.addItem(lab, k)
        self.pri.setCurrentIndex(self.pri.findData(card.get("priority", "med") if card else "med"))
        self.has_date = QCheckBox("Due date"); self.date = QDateEdit(QDate.currentDate()); self.date.setCalendarPopup(True); self.date.setDisplayFormat("yyyy-MM-dd"); self.date.setMinimumWidth(150)
        if card and card.get("dueDate"):
            self.has_date.setChecked(True); self.date.setDate(QDate.fromString(card["dueDate"], "yyyy-MM-dd"))
        self.date.setEnabled(self.has_date.isChecked()); self.has_date.toggled.connect(self.date.setEnabled)
        self.rep = QComboBox()
        for k, lab in REPEATS: self.rep.addItem(("↻ " if k else "") + lab, k)
        self.rep.setCurrentIndex(max(0, self.rep.findData(card.get("repeat", "") if card else "")))
        self.has_until = QCheckBox("until"); self.until = QDateEdit(QDate.currentDate().addMonths(3)); self.until.setCalendarPopup(True); self.until.setDisplayFormat("yyyy-MM-dd"); self.until.setMinimumWidth(150)
        if card and card.get("repeatUntil"): self.has_until.setChecked(True); self.until.setDate(QDate.fromString(card["repeatUntil"], "yyyy-MM-dd"))
        def rep_changed():
            on = bool(self.rep.currentData()); self.has_until.setEnabled(on); self.until.setEnabled(on and self.has_until.isChecked())
            if on and not self.has_date.isChecked(): self.has_date.setChecked(True)
        self.rep.currentIndexChanged.connect(rep_changed); self.has_until.toggled.connect(rep_changed); rep_changed()
        v.addWidget(QLabel("Title *")); v.addWidget(self.title)
        v.addWidget(QLabel("Description")); v.addWidget(self.desc)
        g = QGridLayout(); g.addWidget(QLabel("Column"), 0, 0); g.addWidget(QLabel("Priority"), 0, 1)
        g.addWidget(self.col, 1, 0); g.addWidget(self.pri, 1, 1); v.addLayout(g)
        dr = QHBoxLayout(); dr.addWidget(self.has_date); dr.addWidget(self.date); dr.addSpacing(16); dr.addWidget(QLabel("Repeat")); dr.addWidget(self.rep); dr.addWidget(self.has_until); dr.addWidget(self.until); dr.addStretch(); v.addLayout(dr)
        # labels
        lr = QHBoxLayout(); lr.addWidget(QLabel("Labels")); mb = QPushButton("Manage"); mb.clicked.connect(self.manage_labels); lr.addWidget(mb); lr.addStretch(); v.addLayout(lr)
        self.label_box = QWidget(); self.label_layout = QHBoxLayout(self.label_box); self.label_layout.setContentsMargins(0, 0, 0, 0)
        self.sel_labels = set(card.get("labels", [])) if card else set(); self.rebuild_labels(); v.addWidget(self.label_box)
        # subtasks
        v.addWidget(QLabel("Sub-tasks   (double-click an item to edit it)"))
        self.subs = [dict(s) for s in (card.get("subs", []) if card else [])]
        self.sub_list = QListWidget(); self.sub_list.setFixedHeight(130); self.sub_list.setWordWrap(True); self.sub_list.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        self.sub_list.setStyleSheet("QListWidget::item{padding:3px 6px;margin:0 0 2px 0;border-radius:6px;}QListWidget::item:selected{background:rgba(128,128,128,0.25);color:palette(text);}")
        self.sub_list.itemDoubleClicked.connect(lambda it: self.edit_sub(self.sub_list.row(it))); self.rebuild_subs(); v.addWidget(self.sub_list)
        sr = QHBoxLayout(); self.sub_inp = QLineEdit(); self.sub_inp.setPlaceholderText("Add sub-task… (Enter)")
        self.sub_inp.returnPressed.connect(self.add_sub); ab = QPushButton("+"); ab.clicked.connect(self.add_sub); rb = QPushButton("Remove selected"); rb.clicked.connect(self.remove_sub)
        sr.addWidget(self.sub_inp); sr.addWidget(ab); sr.addWidget(rb); v.addLayout(sr)
        bb = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel)
        bb.button(QDialogButtonBox.Save).setObjectName("primary"); bb.button(QDialogButtonBox.Save).setText("Save changes" if card else "Save task")
        bb.accepted.connect(self.accept); bb.rejected.connect(self.reject); v.addWidget(bb)
        QShortcut(QKeySequence("Ctrl+Return"), self, activated=self.accept)
        if preset:      # from quick-add / calendar / palette
            self.title.setText(preset.get("title", "")); self.pri.setCurrentIndex(self.pri.findData(preset.get("priority", "med")))
            if preset.get("dueDate"): self.has_date.setChecked(True); self.date.setDate(QDate.fromString(preset["dueDate"], "yyyy-MM-dd"))
            self.sel_labels = set(preset.get("labels", [])); self.rebuild_labels()
            if preset.get("repeat"): self.rep.setCurrentIndex(max(0, self.rep.findData(preset["repeat"])))

    def rebuild_labels(self):
        while self.label_layout.count():
            w = self.label_layout.takeAt(0).widget()
            if w: w.deleteLater()
        if not self.store.labels:
            self.label_layout.addWidget(QLabel("No labels yet", objectName="muted"))
        for l in self.store.labels:
            b = QPushButton(l["name"]); b.setCheckable(True); b.setChecked(l["id"] in self.sel_labels)
            b.setStyleSheet(f"QPushButton{{border-color:{l['color']};color:{l['color']};}}QPushButton:checked{{background:{l['color']};color:#fff;}}")
            b.toggled.connect(lambda on, lid=l["id"]: self.sel_labels.add(lid) if on else self.sel_labels.discard(lid))
            self.label_layout.addWidget(b)
        self.label_layout.addStretch()

    def manage_labels(self):
        LabelDialog(self, self.store).exec(); self.rebuild_labels()

    def rebuild_subs(self):
        self.sub_list.clear()
        for s in self.subs:
            self.sub_list.addItem(("✓ " if s.get("done") else "○ ") + s["text"])

    def add_sub(self):
        t = self.sub_inp.text().strip()
        if t:
            self.subs.append({"id": uid(), "text": t, "done": False}); self.sub_inp.clear(); self.rebuild_subs()

    def remove_sub(self):
        r = self.sub_list.currentRow()
        if r >= 0:
            del self.subs[r]; self.rebuild_subs()

    def edit_sub(self, r):
        if not 0 <= r < len(self.subs): return
        t, ok = QInputDialog.getMultiLineText(self, "Edit sub-task", "Text:", self.subs[r]["text"])
        if ok and t.strip(): self.subs[r]["text"] = t.strip(); self.rebuild_subs(); self.sub_list.setCurrentRow(r)

    def accept(self):
        if not self.title.text().strip():
            self.title.setFocus(); return
        super().accept()

    def values(self):
        return dict(title=self.title.text().strip(), desc=self.desc.toPlainText().strip(), col=self.col.currentData(),
                    priority=self.pri.currentData(), dueDate=self.date.date().toString("yyyy-MM-dd") if self.has_date.isChecked() else "",
                    subs=self.subs, labels=list(self.sel_labels), repeat=self.rep.currentData() or "",
                    repeatUntil=self.until.date().toString("yyyy-MM-dd") if self.rep.currentData() and self.has_until.isChecked() else "")


class LabelDialog(QDialog):
    def __init__(self, parent, store):
        super().__init__(parent); self.store = store; self.setWindowTitle("Labels"); self.setMinimumWidth(420)
        self.v = QVBoxLayout(self); self.list_box = QVBoxLayout(); self.v.addLayout(self.list_box)
        r = QHBoxLayout(); self.inp = QLineEdit(); self.inp.setPlaceholderText("New label name"); self.inp.returnPressed.connect(self.add)
        self.color = LABEL_PALETTE[len(store.labels) % len(LABEL_PALETTE)]
        self.cbtn = QPushButton("Colour"); self.cbtn.clicked.connect(self.pick); ab = QPushButton("Add"); ab.setObjectName("primary"); ab.clicked.connect(self.add)
        r.addWidget(self.inp); r.addWidget(self.cbtn); r.addWidget(ab); self.v.addLayout(r)
        cb = QPushButton("Close"); cb.clicked.connect(self.accept); self.v.addWidget(cb)
        self.rebuild()

    def pick(self):
        c = QColorDialog.getColor(QColor(self.color), self, "Label colour")
        if c.isValid():
            self.color = c.name(); self.cbtn.setStyleSheet(f"background:{self.color};color:#fff;")

    def rebuild(self):
        while self.list_box.count():
            w = self.list_box.takeAt(0).widget()
            if w: w.deleteLater()
        if not self.store.labels:
            self.list_box.addWidget(QLabel("No labels yet — add one below, e.g. Client, Urgent, Personal.", objectName="muted"))
        for l in self.store.labels:
            row = QHBoxLayout(); w = QWidget(); w.setLayout(row)
            sw = QPushButton(); sw.setFixedSize(26, 22); sw.setStyleSheet(f"background:{l['color']};border-radius:6px;")
            sw.clicked.connect(lambda _, lab=l: self.recolor(lab))
            name = QLineEdit(l["name"]); name.editingFinished.connect(lambda lab=l, ed=name: self.rename(lab, ed.text()))
            n = self.usage(l["id"]); cnt = QLabel(f"{n} tasks", objectName="muted")
            dl = QPushButton("✕"); dl.setObjectName("ghost"); dl.clicked.connect(lambda _, lab=l: self.delete(lab))
            for x in (sw, name, cnt, dl):
                row.addWidget(x)
            self.list_box.addWidget(w)

    def usage(self, lid):
        n = 0
        for Y in self.store.db["years"].values():
            for kd in Y["kanban"].values():
                for arr in kd.values():
                    n += sum(1 for c in arr if lid in c.get("labels", []))
        return n

    def add(self):
        t = self.inp.text().strip()
        if not t: return
        self.store.labels.append({"id": "l_" + uid(), "name": t, "color": self.color}); self.inp.clear()
        self.color = LABEL_PALETTE[len(self.store.labels) % len(LABEL_PALETTE)]; self.cbtn.setStyleSheet("")
        self.store.save(); self.rebuild()

    def rename(self, lab, t):
        if t.strip() and t.strip() != lab["name"]:
            lab["name"] = t.strip(); self.store.save()

    def recolor(self, lab):
        c = QColorDialog.getColor(QColor(lab["color"]), self, "Label colour")
        if c.isValid():
            lab["color"] = c.name(); self.store.save(); self.rebuild()

    def delete(self, lab):
        if QMessageBox.question(self, "Delete label", f"Delete label \"{lab['name']}\"? It will be removed from {self.usage(lab['id'])} task(s).") != QMessageBox.Yes:
            return
        for Y in self.store.db["years"].values():
            for kd in Y["kanban"].values():
                for arr in kd.values():
                    for c in arr:
                        if "labels" in c: c["labels"] = [x for x in c["labels"] if x != lab["id"]]
        self.store.labels.remove(lab); self.store.save(); self.rebuild()


class ColumnDialog(QDialog):
    def __init__(self, parent, store, col_id=None):
        super().__init__(parent); self.store = store; self.col = store.col(col_id) if col_id else None; self.result_action = None
        self.setWindowTitle("Column settings" if self.col else "New column"); self.setMinimumWidth(400)
        v = QVBoxLayout(self)
        self.name = QLineEdit(self.col["label"] if self.col else ""); self.name.setPlaceholderText("e.g. Review, Blocked, Waiting")
        self.icon = QLineEdit(self.col.get("icon", "") if self.col else "📌"); self.icon.setMaxLength(4); self.icon.setFixedWidth(60)
        self.color = self.col["color"] if self.col else LABEL_PALETTE[len(store.columns) % len(LABEL_PALETTE)]
        self.cbtn = QPushButton("Colour"); self.cbtn.setStyleSheet(f"background:{self.color};color:#fff;"); self.cbtn.clicked.connect(self.pick)
        r = QHBoxLayout(); r.addWidget(self.name); r.addWidget(self.icon); r.addWidget(self.cbtn); v.addWidget(QLabel("Name *")); v.addLayout(r)
        if self.col and self.col["id"] in BUILTIN:
            v.addWidget(QLabel("Built-in column: rename, recolour and move it, but it can't be deleted — the dashboard and month rollover depend on it.", objectName="muted", wordWrap=True))
        br = QHBoxLayout()
        if self.col:
            for txt, act in (("◀ Move left", "left"), ("Move right ▶", "right")):
                b = QPushButton(txt); b.clicked.connect(lambda _, a=act: self.finish(a)); br.addWidget(b)
            if self.col["id"] not in BUILTIN:
                d = QPushButton("Delete"); d.setObjectName("danger"); d.clicked.connect(lambda: self.finish("delete")); br.addWidget(d)
        br.addStretch(); sv = QPushButton("Save"); sv.setObjectName("primary"); sv.clicked.connect(lambda: self.finish("save")); cn = QPushButton("Cancel"); cn.clicked.connect(self.reject)
        br.addWidget(sv); br.addWidget(cn); v.addLayout(br)

    def pick(self):
        c = QColorDialog.getColor(QColor(self.color), self, "Column colour")
        if c.isValid():
            self.color = c.name(); self.cbtn.setStyleSheet(f"background:{self.color};color:#fff;")

    def finish(self, action):
        if action == "save" and not self.name.text().strip():
            self.name.setFocus(); return
        self.result_action = action; self.accept()


class QuickAdd(QDialog):
    """Floating quick-add window: type once, see the parsed priority / labels / due date live.
    Enter saves into the chosen column; Shift+Enter (or 'Full form') opens TaskDialog pre-filled."""
    def __init__(self, win, col_id="todo"):
        super().__init__(win); self.win = win; self.setWindowFlags(Qt.Dialog | Qt.FramelessWindowHint); self.setFixedWidth(640)
        v = QVBoxLayout(self); v.setContentsMargins(12, 12, 12, 10); v.setSpacing(8); top = QHBoxLayout(); top.setSpacing(8)
        self.inp = QLineEdit(); self.inp.setPlaceholderText("What needs to be done?   !high  #label  @fri"); self.inp.textChanged.connect(self.preview); self.inp.returnPressed.connect(self.save)
        self.col = QComboBox()
        for c in win.store.columns: self.col.addItem(f"{c.get('icon', '')} {c['label']}", c["id"])
        self.col.setCurrentIndex(max(0, self.col.findData(col_id))); top.addWidget(self.inp, 1); top.addWidget(self.col); v.addLayout(top)
        self.chips = QHBoxLayout(); self.chips.setSpacing(6); v.addLayout(self.chips)
        foot = QHBoxLayout(); foot.addWidget(QLabel("Enter save · Shift+Enter full form · Esc close     !high !low · #label · @today @tomorrow @fri @15 @+3 @2026-10-01", objectName="muted")); foot.addStretch()
        fb = QPushButton("Full form ▸"); fb.setObjectName("ghost"); fb.clicked.connect(self.full); foot.addWidget(fb); v.addLayout(foot)
        QShortcut(QKeySequence("Shift+Return"), self, activated=self.full); self.preview(); QTimer.singleShot(0, self.inp.setFocus)

    def preview(self, _=None):
        MainWindow._clear(self.chips); s = self.win.store; p = parse_quick(s, self.inp.text(), create=False)
        self.chips.addWidget(badge(PRI_LABEL[p["priority"]], PRI_COLOR[p["priority"]]))
        if p["dueDate"]: self.chips.addWidget(badge("📅 " + fmt_date(p["dueDate"]), "#8890b0"))
        for l in (x for x in s.labels if x["id"] in p["labels"]): self.chips.addWidget(badge(l["name"], "#ffffff", l["color"]))
        for n in p.get("new", []): self.chips.addWidget(badge("+ " + n, "#8890b0"))
        if p.get("repeat"): self.chips.addWidget(badge("↻ " + p["repeat"], "#b48aff"))
        t = QLabel(p["title"] or "Title…", objectName="muted2"); t.setStyleSheet("font-weight:600;"); self.chips.addWidget(t); self.chips.addStretch()

    def save(self):
        if self.win.qa_submit(self.col.currentData(), self.inp.text()): self.accept()

    def full(self):
        col, text = self.col.currentData(), self.inp.text(); self.accept(); self.win.add_task(col, parse_quick(self.win.store, text))


class Palette(QDialog):
    """Ctrl+K: search tasks across every month and year, or run a command."""
    def __init__(self, win):
        super().__init__(win); self.win = win; self.items = []; self.setWindowFlags(Qt.Dialog | Qt.FramelessWindowHint); self.setFixedWidth(640)
        v = QVBoxLayout(self); v.setContentsMargins(10, 10, 10, 8); v.setSpacing(8)
        self.inp = QLineEdit(); self.inp.setPlaceholderText("Search tasks across every month, or type a command…"); self.inp.textChanged.connect(self.render); self.inp.returnPressed.connect(self.run); self.inp.installEventFilter(self)
        self.list = QListWidget(); self.list.setFixedHeight(380); self.list.itemClicked.connect(lambda _: self.run())
        self.list.setStyleSheet(f"QListWidget::item{{padding:6px 8px;margin:0 0 2px 0;border-radius:6px;}}QListWidget::item:selected{{background:{win.palette_dict()['accent']};color:#fff;}}")
        v.addWidget(self.inp); v.addWidget(self.list); v.addWidget(QLabel("↑↓ move  ·  Enter open  ·  Esc close  ·  tasks come from every month and year", objectName="muted"))
        self.render(""); QTimer.singleShot(0, self.inp.setFocus)

    def commands(self):
        w = self.win; light = w.appearance.get("theme") == "light"
        cmds = [("＋", "New task", "Ctrl N", lambda: w.add_task()), ("⚡", "Quick add", "!high #label @fri", lambda: QuickAdd(w).exec()), ("◈", "Go to Dashboard", "Alt 1", lambda: w.show_page("dashboard")), ("⊞", "Go to Task Board", "Alt 2", lambda: w.show_page("kanban")),
                ("☰", "Go to Task List", "Alt 3", lambda: w.show_page("list")), ("▦", "Go to Calendar", "Alt 4", lambda: w.show_page("calendar")),
                ("📍", "Jump to today", f"{MONTHS_LONG[dt.date.today().month - 1]} {dt.date.today().year}", w.go_today),
                ("🌙" if light else "☀️", "Switch to dark theme" if light else "Switch to light theme", "", w.toggle_theme),
                ("⬇", "Export to Excel", "", w.export_excel), ("🔄", "Month rollover", "", w.open_rollover), ("🏷", "Manage labels", "", w.manage_labels), ("🎨", "Appearance", "", w.open_appearance),
                ("💾", "Backup now", "Ctrl Shift B", w.backup_now_ui), ("🗂", "Backup & restore", "", w.open_backup), ("⌨", "Keyboard shortcuts", "?", w.open_shortcuts)]
        cmds += [("📅", f"Open {ml} {w.store.year}", "Switch month", lambda i=i: w.switch_month(i)) for i, ml in enumerate(MONTHS_LONG)]
        return [dict(icon=i, text=t, meta=m, run=r) for i, t, m, r in cmds]

    def tasks(self, q):
        w = self.win; s = w.store; cur = (s.year, s.month); out = []
        def score(title, desc="", subs=()):
            t = title.lower()
            if not q: return 1
            return 4 if t.startswith(q) else 3 if q in t else 2 if q in desc.lower() else 1 if any(q in x.get("text", "").lower() for x in subs) else 0
        for y in s.years():
            for mi in range(12):
                for c in s.columns:
                    for k in s.kanban(mi, y).get(c["id"], []):
                        sc = score(k.get("title", ""), k.get("desc", ""), k.get("subs", []))
                        if sc: out.append(dict(icon="✓" if k.get("done") else "○", text=k["title"], meta=f"{MONTHS[mi]} {y} · {c['label']}", score=sc + 0.5 * ((y, mi) == cur) + 0.2 * (not k.get("done")),
                                              ts=k.get("updatedAt") or k.get("createdAt") or "", run=lambda y=y, mi=mi, kid=k["id"], cid=c["id"]: w.goto(y, mi, "kanban", lambda: w.open_detail(kid, cid))))
        out.sort(key=lambda i: i["ts"], reverse=True); out.sort(key=lambda i: -i["score"]); return out

    def render(self, q):
        raw = q.strip(); q = raw.lower(); self.list.clear(); tasks = self.tasks(q)
        self.items = ([c for c in self.commands() if q in c["text"].lower()] + tasks[:40]) if q else (tasks[:15] + self.commands())
        if q and not self.items: self.items = [dict(icon="＋", text=f"Create task “{raw}”", meta="in To Do", run=lambda t=raw: self.win.qa_submit("todo", t))]
        for it in self.items: self.list.addItem(f"{it['icon']}  {it['text']}" + (f"      —  {it['meta']}" if it["meta"] else ""))
        self.list.setCurrentRow(0)

    def eventFilter(self, obj, e):
        if e.type() == QEvent.KeyPress and e.key() in (Qt.Key_Up, Qt.Key_Down, Qt.Key_PageUp, Qt.Key_PageDown): self.list.keyPressEvent(e); return True
        return super().eventFilter(obj, e)

    def run(self):
        r = self.list.currentRow()
        if 0 <= r < len(self.items): self.accept(); self.items[r]["run"]()


class ExportDialog(QDialog):
    def __init__(self, parent, store):
        super().__init__(parent); self.store = store; self.setWindowTitle("Export to Excel"); self.setMinimumWidth(460)
        v = QVBoxLayout(self); v.addWidget(QLabel(f"Select months to export ({store.year})", objectName="sectitle"))
        g = QGridLayout(); self.checks = []
        for i, m in enumerate(MONTHS):
            cb = QCheckBox(m); cb.setChecked(i == store.month)
            kd = store.kanban(i); has = any(kd.get(c["id"]) for c in store.columns)
            if has: cb.setText(m + " •")
            g.addWidget(cb, i // 4, i % 4); self.checks.append(cb)
        v.addLayout(g)
        r = QHBoxLayout()
        for txt, fn in (("All", lambda: [c.setChecked(True) for c in self.checks]), ("None", lambda: [c.setChecked(False) for c in self.checks]),
                        ("With data", lambda: [c.setChecked("•" in c.text()) for c in self.checks])):
            b = QPushButton(txt); b.clicked.connect(fn); r.addWidget(b)
        r.addStretch(); v.addLayout(r)
        self.inc_board = QCheckBox("Board tasks (with sub-tasks)"); self.inc_board.setChecked(True)
        self.inc_sum = QCheckBox("Month-wise summary sheet"); self.inc_sum.setChecked(True)
        self.inc_done = QCheckBox("Include completed tasks"); self.inc_done.setChecked(True)
        for cb in (self.inc_board, self.inc_sum, self.inc_done): v.addWidget(cb)
        bb = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel); bb.button(QDialogButtonBox.Ok).setText("Export…"); bb.button(QDialogButtonBox.Ok).setObjectName("primary")
        bb.accepted.connect(self.accept); bb.rejected.connect(self.reject); v.addWidget(bb)

    def months(self):
        return [i for i, c in enumerate(self.checks) if c.isChecked()]


def export_excel(store, path, months, inc_board, inc_sum, inc_done):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook(); wb.remove(wb.active)
    hdr = Font(bold=True, color="FFFFFF"); fill = PatternFill("solid", fgColor="2B3563"); title = Font(bold=True, size=13)
    lname = lambda c: ", ".join(l["name"] for l in store.labels if l["id"] in c.get("labels", []))
    if inc_board:
        for mi in months:
            ws = wb.create_sheet(f"{MONTHS[mi]}-Board"); ws.append([f"{APP_NAME} — Board Tasks — {MONTHS[mi]} {store.year}"]); ws["A1"].font = title
            kd = store.kanban(mi)
            for c in store.columns:
                cards = [x for x in kd.get(c["id"], []) if inc_done or not x.get("done")]
                if not cards or (c["id"] == "done" and not inc_done): continue
                ws.append([]); ws.append([f"{c.get('icon', '')} {c['label'].upper()}"]); ws.cell(ws.max_row, 1).font = Font(bold=True, color=c["color"].lstrip("#"))
                ws.append(["#", "Task Title", "Description", "Labels", "Priority", "Due Date", "Repeat", "Status", "Migrated From", "Sub-tasks", "Done/Total", "Created"])
                for cell in ws[ws.max_row]: cell.font = hdr; cell.fill = fill
                for n, x in enumerate(cards, 1):
                    subs = x.get("subs", []); sd = sum(1 for s in subs if s.get("done"))
                    ws.append([n, x["title"], x.get("desc", ""), lname(x), PRI_LABEL[x.get("priority", "med")], x.get("dueDate", ""), x.get("repeat", ""),
                               "Completed" if x.get("done") or c["id"] == "done" else c["label"], x.get("migratedFrom", ""),
                               " | ".join(("[✓] " if s.get("done") else "[ ] ") + s["text"] for s in subs), f"{sd}/{len(subs)}" if subs else "", (x.get("createdAt") or "")[:10]])
            for col, w in zip("ABCDEFGHIJKL", (4, 30, 28, 16, 10, 12, 10, 12, 13, 36, 10, 12)): ws.column_dimensions[col].width = w
    if inc_sum:
        ws = wb.create_sheet("Summary"); ws.append([f"{APP_NAME} — Month-wise Summary — {store.year}"]); ws["A1"].font = title; ws.append([])
        ws.append(["Month", "Year", "To Do", "Today's", "In Progress", "Completed", "Total", "Completion %"])
        for cell in ws[ws.max_row]: cell.font = hdr; cell.fill = fill
        ip = [c["id"] for c in store.columns if c["id"] not in ("todo", "today", "done")]
        for mi in range(12):
            kd = store.kanban(mi); todo, today, done = len(kd.get("todo", [])), len(kd.get("today", [])), len(kd.get("done", []))
            wip = sum(len(kd.get(i, [])) for i in ip); tot = todo + today + wip + done
            ws.append([MONTHS[mi], store.year, todo, today, wip, done, tot, f"{round(done / tot * 100) if tot else 0}%"])
    if not wb.sheetnames: wb.create_sheet("Empty").append(["Nothing selected"])
    wb.save(path)


# ═══════════════════════════════════════════════════════════════════════════
#  MAIN WINDOW
# ═══════════════════════════════════════════════════════════════════════════
class MainWindow(QMainWindow):
    def __init__(self, store):
        super().__init__(); self.store = store; self.setWindowTitle(APP_NAME); self.resize(1400, 880); self.setMinimumSize(980, 620)
        self.detail_ref = None; self.filter_q = ""; self.filter_pri = ""; self.filter_labels = set()
        self.settings = read_settings(); self.appearance = {**dict(theme="dark", font="", size=13, text="", accent="", bg="", surface="", motion=True), **self.settings.get("appearance", {}), **self.store.db["meta"].get("appearance_py", {})}
        self._build(); self.apply_appearance(); self._tray(); self.show_page("dashboard")
        if store.merged_count: QTimer.singleShot(800, lambda: self.toast(f"Checklist merged into the board: {store.merged_count} tasks now carry their group as a label. A backup of the old file was saved first.", 7000))
        moved = self.store.auto_migrate()
        if moved:
            self.toast(f"🔄 {moved} incomplete task(s) auto-moved to the next month"); self.refresh()
        self.backup_timer = QTimer(self); self.backup_timer.timeout.connect(lambda: self.store.backup_now()); self.backup_timer.start(30 * 60 * 1000)
        QShortcut(QKeySequence("Ctrl+N"), self, activated=lambda: self.add_task())
        QShortcut(QKeySequence("Ctrl+Shift+B"), self, activated=self.backup_now_ui)
        QShortcut(QKeySequence("Ctrl+K"), self, activated=lambda: Palette(self).exec())
        QShortcut(QKeySequence("["), self, activated=lambda: self.step_month(-1)); QShortcut(QKeySequence("]"), self, activated=lambda: self.step_month(1))
        QShortcut(QKeySequence("?"), self, activated=self.open_shortcuts)
        for i, k in enumerate(("dashboard", "kanban", "list", "calendar")): QShortcut(QKeySequence(f"Alt+{i + 1}"), self, activated=lambda k=k: self.show_page(k))

    # ---------- layout ----------
    def _build(self):
        root = QWidget(); self.setCentralWidget(root); h = QHBoxLayout(root); h.setContentsMargins(0, 0, 0, 0); h.setSpacing(0)
        sb = QFrame(); sb.setObjectName("sidebar"); sb.setFixedWidth(220); sv = QVBoxLayout(sb); sv.setContentsMargins(0, 20, 0, 16); sv.setSpacing(2)
        b = QLabel(APP_NAME); b.setObjectName("brand"); b.setContentsMargins(20, 0, 0, 0); sv.addWidget(b)
        bs = QLabel("PLAN · DO · TRACK"); bs.setObjectName("brandSub"); bs.setContentsMargins(20, 0, 0, 14); sv.addWidget(bs)
        sv.addWidget(QLabel("WORKSPACE", objectName="navSection", indent=20)); self.nav_btns = {}
        for key, txt in (("dashboard", "◈  Dashboard"), ("kanban", "⊞  Task Board"), ("list", "☰  Task List"), ("calendar", "▦  Calendar")):
            nb = QPushButton(txt); nb.setObjectName("nav"); nb.setCheckable(True); nb.clicked.connect(lambda _, k=key: self.show_page(k)); sv.addWidget(nb); self.nav_btns[key] = nb
        sv.addSpacing(10); sv.addWidget(QLabel("TOOLS", objectName="navSection", indent=20))
        self.roll_btn = QPushButton("🔄  Month Rollover"); self.roll_btn.setObjectName("nav"); self.roll_btn.clicked.connect(self.open_rollover); sv.addWidget(self.roll_btn)
        for txt, fn in (("💾  Backup && Restore", self.open_backup), ("🎨  Appearance", self.open_appearance), ("⌕  Search   ·   Ctrl K", lambda: Palette(self).exec())):
            nb = QPushButton(txt); nb.setObjectName("nav"); nb.clicked.connect(fn); sv.addWidget(nb)
        sv.addStretch()
        yr = QHBoxLayout(); yr.setContentsMargins(16, 0, 16, 6)
        pb = QPushButton("◀"); pb.setObjectName("ghost"); pb.clicked.connect(lambda: self.switch_year(-1)); self.year_lbl = QLabel(str(self.store.year)); self.year_lbl.setAlignment(Qt.AlignCenter); self.year_lbl.setStyleSheet("font-weight:700;font-size:15px;")
        nb2 = QPushButton("▶"); nb2.setObjectName("ghost"); nb2.clicked.connect(lambda: self.switch_year(1)); yr.addWidget(pb); yr.addWidget(self.year_lbl, 1); yr.addWidget(nb2); sv.addLayout(yr)
        mg = QGridLayout(); mg.setContentsMargins(16, 0, 16, 0); mg.setSpacing(4); self.month_btns = []
        for i, m in enumerate(MONTHS):
            mb = QPushButton(m); mb.setObjectName("month"); mb.setCheckable(True); mb.clicked.connect(lambda _, i=i: self.switch_month(i)); mg.addWidget(mb, i // 3, i % 3); self.month_btns.append(mb)
        sv.addLayout(mg); h.addWidget(sb)
        main = QVBoxLayout(); main.setContentsMargins(0, 0, 0, 0); main.setSpacing(0)
        tb = QFrame(); tb.setObjectName("topbar"); th = QHBoxLayout(tb); th.setContentsMargins(24, 12, 24, 12)
        self.page_title = QLabel("Dashboard"); self.page_title.setObjectName("pageTitle"); self.page_month = QLabel(); self.page_month.setObjectName("pageMonth"); self.page_month.setMargin(6)
        th.addWidget(self.page_title); th.addWidget(self.page_month); th.addStretch()
        self.theme_btn = QPushButton("🌙"); self.theme_btn.setToolTip("Toggle dark/light"); self.theme_btn.clicked.connect(self.toggle_theme)
        ex = QPushButton("⬇ Export Excel"); ex.clicked.connect(self.export_excel); ad = QPushButton("+ Add Task"); ad.setObjectName("primary"); ad.setToolTip("New task (Ctrl N)"); ad.clicked.connect(lambda: self.add_task())
        th.addWidget(self.theme_btn); th.addWidget(ex); th.addWidget(ad); main.addWidget(tb)
        self.stack = QStackedWidget(); main.addWidget(self.stack, 1)
        self.pages = {}
        for key, build in (("dashboard", self._build_dashboard), ("kanban", self._build_board), ("list", self._build_list), ("calendar", self._build_calendar)):
            self.pages[key] = build(); self.stack.addWidget(self.pages[key])
        w = QWidget(); w.setLayout(main); h.addWidget(w, 1)
        # detail dock
        self.dock = QDockWidget("Task", self); self.dock.setAllowedAreas(Qt.RightDockWidgetArea); self.dock.setFeatures(QDockWidget.DockWidgetClosable); self.dock.setMinimumWidth(440)
        self.detail = QFrame(); self.detail.setObjectName("detail"); self.detail_layout = QVBoxLayout(self.detail); self.dock.setWidget(self.detail); self.addDockWidget(Qt.RightDockWidgetArea, self.dock); self.dock.hide()
        self.dock.visibilityChanged.connect(lambda vis: setattr(self, "detail_ref", None) if not vis else None)
        self.toast_lbl = QLabel(self); self.toast_lbl.setStyleSheet("background:#13161f;color:#eef0f8;border:1px solid #3ddbbf;border-radius:10px;padding:8px 16px;"); self.toast_lbl.hide()

    def _scroll(self, inner):
        sa = QScrollArea(); sa.setWidgetResizable(True); sa.setWidget(inner); return sa

    def _build_dashboard(self):
        w = QWidget(); v = QVBoxLayout(w); v.setContentsMargins(24, 20, 24, 20); v.setSpacing(16)
        self.kpi_row = QHBoxLayout(); v.addLayout(self.kpi_row)
        row = QHBoxLayout(); v.addLayout(row, 1)
        p0 = QFrame(); p0.setObjectName("panel"); l0 = QVBoxLayout(p0); hr = QHBoxLayout(); hr.addWidget(QLabel("Needs attention", objectName="sectitle")); self.attn_sub = QLabel(objectName="muted"); hr.addStretch(); hr.addWidget(self.attn_sub); l0.addLayout(hr)
        self.attn_box = QVBoxLayout(); l0.addLayout(self.attn_box); l0.addStretch()
        l0.addWidget(QLabel("Open tasks by priority", objectName="muted")); self.pri_mix = PriMix(); l0.addWidget(self.pri_mix); self.pri_legend = QLabel(objectName="muted"); l0.addWidget(self.pri_legend); row.addWidget(p0, 3)
        p1 = QFrame(); p1.setObjectName("panel"); l1 = QVBoxLayout(p1); l1.addWidget(QLabel("Recent Activity", objectName="sectitle")); self.activity_box = QVBoxLayout(); l1.addLayout(self.activity_box); l1.addStretch(); row.addWidget(p1, 3)
        p2 = QFrame(); p2.setObjectName("panel"); l2 = QVBoxLayout(p2); l2.addWidget(QLabel("Tasks / Month", objectName="sectitle"))
        self.chart = BarChart(lambda: [sum(len(self.store.kanban(i).get(c["id"], [])) for c in self.store.columns) for i in range(12)], "#6c8aff"); self.chart.clicked.connect(self.switch_month); l2.addWidget(self.chart); l2.addWidget(QLabel("Click a bar to open that month", objectName="muted")); l2.addStretch(); row.addWidget(p2, 2)
        return self._scroll(w)

    def _build_board(self):
        w = QWidget(); v = QVBoxLayout(w); v.setContentsMargins(24, 14, 24, 14); v.setSpacing(12)
        tbar = QHBoxLayout(); self.search = QLineEdit(); self.search.setPlaceholderText("Search tasks…"); self.search.setMinimumWidth(160); self.search.setMaximumWidth(260); self.search.textChanged.connect(self.on_search)
        self.pri_filter = QComboBox(); self.pri_filter.setMinimumWidth(140); self.pri_filter.addItem("All priorities", "")
        for k, lab in PRI_LABEL.items(): self.pri_filter.addItem(lab, k)
        self.pri_filter.currentIndexChanged.connect(lambda: (setattr(self, "filter_pri", self.pri_filter.currentData()), self.render_board()))
        self.label_bar = QHBoxLayout(); tbar.addWidget(self.search); tbar.addWidget(self.pri_filter); tbar.addLayout(self.label_bar); tbar.addStretch()
        lb = QPushButton("🏷 Labels"); lb.clicked.connect(self.manage_labels); cb = QPushButton("+ Column"); cb.clicked.connect(lambda: self.column_dialog(None)); tbar.addWidget(lb); tbar.addWidget(cb); v.addLayout(tbar)
        self.board_host = QWidget(); self.board_layout = QHBoxLayout(self.board_host); self.board_layout.setContentsMargins(0, 0, 0, 0); self.board_layout.setSpacing(12); self.board_layout.setAlignment(Qt.AlignLeft | Qt.AlignTop)
        sa = QScrollArea(); sa.setWidgetResizable(True); sa.setWidget(self.board_host); v.addWidget(sa, 1)
        return w

    def _build_list(self):
        w = QWidget(); v = QVBoxLayout(w); v.setContentsMargins(24, 14, 24, 14); v.setSpacing(12); tb = QHBoxLayout()
        tb.addWidget(QLabel("Group by", objectName="muted")); self.list_group = QComboBox()
        for k, t in (("column", "Column"), ("due", "Due date"), ("label", "Label"), ("priority", "Priority")): self.list_group.addItem(t, k)
        self.list_group.currentIndexChanged.connect(self.render_list); tb.addWidget(self.list_group)
        self.list_search = QLineEdit(); self.list_search.setPlaceholderText("Search tasks…"); self.list_search.setFixedWidth(240); self.list_search.textChanged.connect(self.render_list); tb.addWidget(self.list_search)
        self.list_hide_done = QCheckBox("Hide completed"); self.list_hide_done.toggled.connect(self.render_list); tb.addWidget(self.list_hide_done)
        self.list_stat = QLabel(objectName="muted"); tb.addWidget(self.list_stat); tb.addStretch(); tb.addWidget(QLabel("Same tasks as the board · tick = done · click a row to open", objectName="muted")); v.addLayout(tb)
        self.list_host = QWidget(); self.list_layout = QVBoxLayout(self.list_host); self.list_layout.setAlignment(Qt.AlignTop); self.list_layout.setSpacing(6); v.addWidget(self._scroll(self.list_host), 1); return w

    def _build_calendar(self):
        w = QWidget(); v = QVBoxLayout(w); v.setContentsMargins(24, 14, 24, 14); v.setSpacing(12); tb = QHBoxLayout()
        for t, fn, tip in (("‹", lambda: self.step_month(-1), "Previous month  [ "), ("›", lambda: self.step_month(1), "Next month  ] "), ("Today", self.go_today, "Jump to today")):
            b = QPushButton(t); b.setToolTip(tip); b.clicked.connect(fn); tb.addWidget(b)
        self.cal_stat = QLabel(objectName="muted"); tb.addWidget(self.cal_stat); tb.addStretch()
        tb.addWidget(QLabel("Drag to reschedule · double-click a day to add", objectName="muted")); v.addLayout(tb)
        row = QHBoxLayout(); row.setSpacing(14); gw = QWidget(); self.cal_grid = QGridLayout(gw); self.cal_grid.setSpacing(6); self.cal_grid.setAlignment(Qt.AlignTop); row.addWidget(self._scroll(gw), 1)
        sw = QWidget(); self.cal_side = QVBoxLayout(sw); self.cal_side.setContentsMargins(0, 0, 0, 0); self.cal_side.setAlignment(Qt.AlignTop); ss = self._scroll(sw); ss.setFixedWidth(250); row.addWidget(ss); v.addLayout(row, 1); return w

    # ---------- tray ----------
    def _tray(self):
        pm = QPixmap(64, 64); pm.fill(Qt.transparent); p = QPainter(pm); p.setRenderHint(QPainter.Antialiasing); p.setBrush(QColor("#6c8aff")); p.setPen(Qt.NoPen); p.drawRoundedRect(4, 4, 56, 56, 14, 14)
        p.setPen(QColor("#ffffff")); f = QFont(); f.setBold(True); f.setPixelSize(34); p.setFont(f); p.drawText(pm.rect(), Qt.AlignCenter, "T"); p.end()
        self.setWindowIcon(QIcon(pm)); self.tray = QSystemTrayIcon(QIcon(pm), self); m = QMenu()
        m.addAction(f"Open {APP_NAME}", self.show_window); m.addSeparator(); m.addAction("Backup Data Now", self.backup_now_ui)
        m.addAction("Open Backup Folder", lambda: self.open_path(get_backup_dir())); m.addAction("Open Data Folder", lambda: self.open_path(USER_DATA)); m.addSeparator(); m.addAction("Quit", self.quit_app)
        self.tray.setContextMenu(m); self.tray.activated.connect(lambda r: self.show_window() if r == QSystemTrayIcon.DoubleClick else None); self.tray.setToolTip(APP_NAME); self.tray.show()
        self.quitting = False; self._last_close = 0; self._hint_shown = False

    def show_window(self):
        self.show(); self.raise_(); self.activateWindow()

    def quit_app(self):
        self.quitting = True; self.store.save(); self.store.backup_now(); QApplication.quit()

    def closeEvent(self, e):
        if self.quitting: e.accept(); return
        if time.time() - self._last_close < 3: self.quitting = True; e.accept(); QApplication.quit(); return
        self._last_close = time.time(); e.ignore(); self.hide()
        if not self._hint_shown:
            self._hint_shown = True; self.tray.showMessage(APP_NAME, "Still running in the system tray. Right-click the tray icon and choose Quit to exit.")

    @staticmethod
    def open_path(p):
        try:
            if sys.platform.startswith("win"): os.startfile(p)
            elif sys.platform == "darwin": os.system(f'open "{p}"')
            else: os.system(f'xdg-open "{p}" &')
        except Exception:
            pass

    # ---------- helpers ----------
    def toast(self, msg, ms=2600):
        self.toast_lbl.setText(msg); self.toast_lbl.adjustSize(); self._place_toast(); self.toast_lbl.show(); self.toast_lbl.raise_()
        if self.motion():       # slide up from just below its resting spot
            end = self.toast_lbl.pos(); self.animate(self.toast_lbl, b"pos", end + QPoint(0, 18), end, 220)
        QTimer.singleShot(ms, self.toast_lbl.hide)

    def motion(self):
        return bool(self.appearance.get("motion", True))

    def animate(self, target, prop, start, end, ms=260, easing=QEasingCurve.OutCubic, on_value=None, on_done=None):
        """One helper for every transition: QVariantAnimation calling on_value(v), or QPropertyAnimation on `prop`.
        The animation is a child of `target`, so a re-render that deletes the widget also stops the animation."""
        for old in target.findChildren(QAbstractAnimation): old.stop()      # one animation per target at a time
        if not self.motion():       # jump straight to the final state
            on_value(end) if on_value else target.setProperty(prop, end)
            if on_done: on_done()
            return None
        a = QVariantAnimation(target) if on_value else QPropertyAnimation(target, prop, target)
        a.setStartValue(start); a.setEndValue(end); a.setDuration(ms); a.setEasingCurve(easing)
        if on_value: a.valueChanged.connect(on_value)
        if on_done: a.finished.connect(on_done)
        a.start(QAbstractAnimation.DeleteWhenStopped); return a

    def fade_in(self, w, ms=180):
        if not self.motion(): return
        eff = QGraphicsOpacityEffect(w); w.setGraphicsEffect(eff)
        self.animate(eff, b"opacity", 0.0, 1.0, ms, QEasingCurve.OutQuad, on_done=lambda: w.setGraphicsEffect(None))

    def resizeEvent(self, e):
        super().resizeEvent(e)
        if self.toast_lbl.isVisible(): self._place_toast()

    def _place_toast(self):
        c = self.centralWidget().geometry()
        self.toast_lbl.move(c.right() - self.toast_lbl.width() - 24, c.bottom() - self.toast_lbl.height() - 24)

    def show_page(self, key):
        changed = getattr(self, "page", None) != key; self.page = key
        for k, b in self.nav_btns.items(): b.setChecked(k == key)
        self.stack.setCurrentWidget(self.pages[key]); self.page_title.setText({"dashboard": "Dashboard", "kanban": "Task Board", "list": "Task List", "calendar": "Calendar"}[key]); self.refresh()
        if changed: self.fade_in(self.pages[key])

    def refresh(self):
        s = self.store; self.page_month.setText(f"{MONTHS_LONG[s.month]} {s.year}"); self.year_lbl.setText(str(s.year))
        for i, b in enumerate(self.month_btns): b.setChecked(i == s.month)
        kd = s.kanban(); n = sum(len(kd.get(c["id"], [])) for c in s.columns); nopen = sum(1 for c in s.columns if c["id"] != "done" for x in kd.get(c["id"], []) if not x.get("done"))
        self.nav_btns["kanban"].setText(f"⊞  Task Board   ({n})"); self.nav_btns["list"].setText(f"☰  Task List   ({nopen} open)")
        self.overdue_n = sum(1 for c in s.columns if c["id"] != "done" for x in kd.get(c["id"], []) if is_overdue(x.get("dueDate"), x.get("done")))
        self.nav_btns["calendar"].setText("▦  Calendar" + (f"   (⚠ {self.overdue_n})" if self.overdue_n else ""))
        self.roll_btn.setText("🔄  Month Rollover" + ("  ⚠" if s.pending_rollover() else ""))
        {"dashboard": self.render_dashboard, "kanban": self.render_board, "list": self.render_list, "calendar": self.render_calendar}[self.page]()
        if self.detail_ref: self.render_detail()

    def switch_month(self, i):
        self.store.month = i; self.close_detail(); self.refresh()

    def switch_year(self, d):
        ys = self.store.years(); ry = dt.date.today().year; lo, hi = min(ys[0] if ys else ry, ry - 1), max(ys[-1] if ys else ry, ry + 1)
        ny = self.store.year + d
        if lo <= ny <= hi: self.store.year = ny; self.store.ensure_year(ny); self.close_detail(); self.refresh()
        else: self.toast(f"Years available: {lo} – {hi}")

    def step_month(self, d):
        nm = self.store.month + d
        if 0 <= nm <= 11: self.switch_month(nm); return
        before = self.store.year; self.switch_year(d)
        if self.store.year != before: self.switch_month(11 if nm < 0 else 0)

    def go_today(self):
        t = dt.date.today(); self.store.year = t.year; self.store.ensure_year(t.year); self.switch_month(t.month - 1)

    def goto(self, year, mi, page, then=None):
        """Palette: jump to a task's month/year, open its page, then run `then` (e.g. open the detail panel)."""
        self.store.year = year; self.store.ensure_year(year); self.store.month = mi; self.close_detail(); self.show_page(page)
        if then: then()

    def open_shortcuts(self):
        QMessageBox.information(self, "Keyboard shortcuts", "Ctrl K\tSearch & commands\nCtrl N\tNew task\n[  ]\tPrevious / next month\nAlt 1–4\tDashboard · Board · List · Calendar\n"
                                "Enter\tQuick-add: save task\nShift Enter\tQuick-add: open full form\nCtrl Enter\tSave the open form\n?\tThis sheet\nCtrl Shift B\tBackup now\n\n"
                                "Quick-add syntax:  !high / !low priority  ·  #label assigns (or creates) a label  ·  @today @tomorrow @fri @15 @+3 @2026-10-01 due date  ·  *daily *weekdays *weekly *monthly *yearly repeat")

    @staticmethod
    def _clear(layout):
        while layout.count():
            it = layout.takeAt(0)
            if it.widget(): it.widget().deleteLater()
            elif it.layout(): MainWindow._clear(it.layout())

    # ---------- dashboard ----------
    def render_dashboard(self):
        s = self.store; kd = s.kanban(); self._clear(self.kpi_row)
        todo, today, done = len(kd.get("todo", [])), len(kd.get("today", [])), len(kd.get("done", []))
        wip = sum(len(kd.get(c["id"], [])) for c in s.columns if c["id"] not in ("todo", "today", "done")); total = todo + today + wip + done
        pct = round(done / total * 100) if total else 0
        overdue = self.overdue_n
        for lab, val, sub, color in (("TOTAL TASKS", total, f"{overdue} overdue" if overdue else "This month", "#6c8aff"), ("TODAY'S FOCUS", today, "Tasks for today", "#ffb347"),
                                     ("IN PROGRESS", wip, "Currently working", "#b48aff"), ("COMPLETED", done, f"{pct}% of month", "#3ddbbf")):
            f = QFrame(); f.setObjectName("kpi"); fh = QHBoxLayout(f); fv = QVBoxLayout(); fv.addWidget(QLabel(lab, objectName="kpiLabel")); vl = QLabel("0", objectName="kpiVal"); fv.addWidget(vl)
            sl = QLabel(sub); sl.setObjectName("muted"); sl.setStyleSheet(f"color:{'#ff6584' if 'overdue' in sub else '#5c6382'}"); fv.addWidget(sl); fh.addLayout(fv, 1); f.setStyleSheet(f"QFrame#kpi{{border-bottom:3px solid {color};}}"); self.kpi_row.addWidget(f)
            self.animate(vl, None, 0, val, 520, on_value=lambda n, l=vl: l.setText(str(int(n))))
            if lab == "COMPLETED":
                ring = Ring(color); fh.addWidget(ring, 0, Qt.AlignVCenter); self.animate(ring, None, 0.0, float(pct), 700, on_value=lambda x, r=ring: (setattr(r, "value", x), r.update()))
        self._clear(self.attn_box); items = []; today = dt.date.today()
        def days_to(iso):
            try: return (dt.date.fromisoformat(iso) - today).days
            except Exception: return None
        for c in s.columns:
            for k in kd.get(c["id"], []) if c["id"] != "done" else []:
                df = days_to(k["dueDate"]) if k.get("dueDate") and not k.get("done") else None
                if df is not None and df <= 7: items.append((df, k.get("priority") != "high", k["title"], c["color"], lambda cid=k["id"], col=c["id"]: (self.show_page("kanban"), self.open_detail(cid, col))))
        items.sort(key=lambda x: x[:2]); over = sum(1 for i in items if i[0] < 0)
        self.attn_sub.setText(f"{over} overdue · {len(items) - over} due this week" if items else "")
        if not items: self.attn_box.addWidget(QLabel("Nothing is overdue or due in the next 7 days. Add due dates to tasks and they show up here.", objectName="muted", wordWrap=True))
        for df, _, title, color, fn in items[:8]:
            when = f"{-df}d overdue" if df < 0 else "Today" if df == 0 else "Tomorrow" if df == 1 else (today + dt.timedelta(df)).strftime("%a") if df <= 6 else fmt_date((today + dt.timedelta(df)).isoformat())
            b = QPushButton(f"{title}     {when}"); b.setObjectName("addk"); b.setStyleSheet(f"QPushButton#addk{{border-left:3px solid {color};" + ("color:#ff6584;" if df < 0 else "") + "}"); b.clicked.connect(lambda _, f=fn: f()); self.attn_box.addWidget(b)
        if len(items) > 8: self.attn_box.addWidget(QLabel(f"+{len(items) - 8} more in the calendar", objectName="muted"))
        open_ = [x for c in s.columns if c["id"] != "done" for x in kd.get(c["id"], [])]; n = {k: sum(1 for x in open_ if x.get("priority", "med") == k) for k in ("high", "med", "low")}
        self.pri_mix.counts = (n["high"], n["med"], n["low"]); self.pri_mix.update()
        self.pri_legend.setText(f"<b style='color:{PRI_COLOR['high']}'>{n['high']}</b> high &nbsp; <b style='color:{PRI_COLOR['med']}'>{n['med']}</b> medium &nbsp; <b style='color:{PRI_COLOR['low']}'>{n['low']}</b> low" if open_ else "No open tasks")
        self._clear(self.activity_box)
        cards = [(x, c) for c in s.columns for x in kd.get(c["id"], [])]
        cards.sort(key=lambda t: t[0].get("updatedAt") or t[0].get("createdAt") or "", reverse=True)
        if not cards: self.activity_box.addWidget(QLabel("No tasks this month yet.", objectName="muted"))
        for x, c in cards[:6]:
            b = QPushButton(("✓ " if x.get("done") else "○ ") + x["title"] + f"   ·  {c['label']} · {PRI_LABEL[x.get('priority', 'med')]}"); b.setObjectName("addk")
            b.clicked.connect(lambda _, cid=x["id"], col=c["id"]: (self.show_page("kanban"), self.open_detail(cid, col))); self.activity_box.addWidget(b)
        self.chart.current = s.month; self.animate(self.chart, None, 0.0, 1.0, 600, on_value=lambda x: (setattr(self.chart, "progress", x), self.chart.update()))

    # ---------- board ----------
    def on_search(self, t):
        self.filter_q = t.lower(); self.render_board()

    def card_matches(self, c):
        if self.filter_pri and c.get("priority") != self.filter_pri: return False
        if self.filter_labels and not self.filter_labels <= set(c.get("labels", [])): return False
        if self.filter_q:
            hay = " ".join([c.get("title", ""), c.get("desc", "")] + [s["text"] for s in c.get("subs", [])]).lower()
            if self.filter_q not in hay: return False
        return True

    def render_board(self):
        s = self.store; self._clear(self.label_bar)
        for l in s.labels:
            b = QPushButton(l["name"]); b.setCheckable(True); b.setChecked(l["id"] in self.filter_labels)
            b.setStyleSheet(f"QPushButton{{border-color:{l['color']};color:{l['color']};padding:3px 10px;}}QPushButton:checked{{background:{l['color']};color:#fff;}}")
            b.setSizePolicy(QSizePolicy.Fixed, QSizePolicy.Fixed)
            b.toggled.connect(lambda on, lid=l["id"]: (self.filter_labels.add(lid) if on else self.filter_labels.discard(lid), self.render_board())); self.label_bar.addWidget(b)
        self._clear(self.board_layout); kd = s.kanban()
        for col in s.columns:
            cw = QWidget(); cw.setMinimumWidth(176); cw.setMaximumWidth(380); cv = QVBoxLayout(cw); cv.setContentsMargins(0, 0, 0, 0); cv.setSpacing(8)
            head = QFrame(); head.setObjectName("colhead"); hh = QHBoxLayout(head); hh.setContentsMargins(12, 8, 8, 8)
            dot = QLabel("●"); dot.setStyleSheet(f"color:{col['color']};"); t = QLabel(f"{col.get('icon', '')} {col['label']}"); t.setStyleSheet("font-weight:700;")
            pal = self.palette_dict(); cnt = badge(str(len(kd.get(col["id"], []))), pal["text2"], pal["s4"]); mb = QPushButton("…"); mb.setObjectName("ghost"); mb.setFixedWidth(26); mb.clicked.connect(lambda _, cid=col["id"]: self.column_dialog(cid))
            hh.addWidget(dot); hh.addWidget(t); hh.addStretch(); hh.addWidget(cnt); hh.addWidget(mb); cv.addWidget(head)
            lst = CardList(col["id"]); lst.dropped.connect(self.on_drop)
            cards = [c for c in kd.get(col["id"], []) if self.card_matches(c)]
            if not kd.get(col["id"]):
                it = QListWidgetItem("📭  Empty"); it.setFlags(Qt.NoItemFlags); it.setTextAlignment(Qt.AlignCenter); lst.addItem(it)
            elif not cards:
                it = QListWidgetItem("🔍  No matches"); it.setFlags(Qt.NoItemFlags); it.setTextAlignment(Qt.AlignCenter); lst.addItem(it)
            for c in cards:
                it = QListWidgetItem(); it.setData(Qt.UserRole, c["id"]); w = CardWidget(s, c, col); it.setSizeHint(w.sizeHint() + QSize(0, 6)); lst.addItem(it); lst.setItemWidget(it, w)
                w.clicked.connect(self.open_detail); w.edit.connect(self.edit_task); w.delete.connect(self.delete_task); w.toggle.connect(self.toggle_done)
            lst.fit_items(); cv.addWidget(lst, 1)
            ab = QPushButton("+ Add Task"); ab.setObjectName("addk"); ab.setToolTip("New task in this column  ·  Ctrl K → Quick add for the !high #label @fri syntax"); ab.clicked.connect(lambda _, cid=col["id"]: self.add_task(cid)); cv.addWidget(ab); self.board_layout.addWidget(cw, 1)
        addc = QPushButton("+"); addc.setObjectName("addk"); addc.setFixedWidth(44); addc.setToolTip("Add column"); addc.clicked.connect(lambda: self.column_dialog(None)); self.board_layout.addWidget(addc, 0, Qt.AlignTop)

    def on_drop(self, cid, from_col, to_col, before):
        if before == cid: return
        card = self.store.move_card(cid, from_col, to_col, before)
        if card:
            self.store.save(); self.refresh()
            if self.detail_ref and self.detail_ref[0] == cid: self.detail_ref = (cid, to_col); self.render_detail()
            if card.pop("recurred", False): self.toast(f"↻ Done for now — next due {fmt_date(card['dueDate'])}")
            elif to_col == "done" and from_col != "done": self.toast("✅ Task completed!")

    def _create_card(self, col_id, title, priority="med", dueDate="", labels=None, desc="", subs=None, repeat="", repeatUntil=""):
        s = self.store; card = dict(id=uid(), title=title, desc=desc, priority=priority, dueDate=dueDate, done=col_id == "done", subs=subs or [], labels=labels or [],
                                    repeat=repeat, repeatUntil=repeatUntil, comments=[], activity=[], createdAt=now_iso(), month=MONTHS[s.month], year=s.year)
        s.log(card, f"Created in {s.col(col_id)['label']}"); s.kanban().setdefault(col_id, []).insert(0, card); s.save(); self.refresh(); return card

    def add_task(self, col_id="todo", preset=None):
        d = TaskDialog(self, self.store, None, col_id, preset)
        if d.exec() != QDialog.Accepted: return
        v = d.values(); self._create_card(v.pop("col"), **v); self.toast("Task added!")

    def qa_submit(self, col_id, text):
        """Quick-add: 'Renew SSL !high #ops @fri' → new card in col_id (returns it, or None if there's no title)."""
        p = parse_quick(self.store, text)
        if not p["title"]: return None
        card = self._create_card(col_id, **p); self.toast(f"Added “{p['title']}”" + (f" · due {fmt_date(p['dueDate'])}" if p["dueDate"] else "")); return card

    def edit_task(self, cid, col_id):
        card, col_id = self.store.find_card(cid, col_id)
        if not card: return
        d = TaskDialog(self, self.store, card, col_id)
        if d.exec() != QDialog.Accepted: return
        v = d.values(); changes = [n for n, k in (("title", "title"), ("description", "desc"), ("priority", "priority"), ("due date", "dueDate"), ("repeat", "repeat")) if (card.get(k) or "") != v[k]]
        card.update(title=v["title"], desc=v["desc"], priority=v["priority"], dueDate=v["dueDate"], subs=v["subs"], labels=v["labels"], repeat=v["repeat"], repeatUntil=v["repeatUntil"], updatedAt=now_iso())
        if changes: self.store.log(card, "Edited " + ", ".join(changes))
        if v["col"] != col_id: self.store.move_card(cid, col_id, v["col"])
        elif v["col"] == "done": card["done"] = True
        if self.detail_ref and self.detail_ref[0] == cid: self.detail_ref = (cid, v["col"])
        self.store.save(); self.refresh(); self.toast("Task updated")

    def delete_task(self, cid, col_id):
        if QMessageBox.question(self, "Delete task", "Delete this task? This cannot be undone.") != QMessageBox.Yes: return
        kd = self.store.kanban(); kd[col_id] = [c for c in kd.get(col_id, []) if c["id"] != cid]
        if self.detail_ref and self.detail_ref[0] == cid: self.close_detail()
        self.store.save(); self.refresh()

    def toggle_done(self, cid, col_id):
        to = "done" if col_id != "done" else "todo"; card = self.store.move_card(cid, col_id, to)
        if card:
            self.store.save(); rec = card.pop("recurred", False)
            if self.detail_ref and self.detail_ref[0] == cid: self.detail_ref = (cid, col_id if rec else to)
            self.refresh(); self.toast(f"↻ Done for now — next due {fmt_date(card['dueDate'])}" if rec else "✅ Moved to Completed!" if to == "done" else "↩ Moved to To Do")

    def skip_occurrence(self, cid, col_id):
        card, _ = self.store.find_card(cid, col_id)
        if not card or not card.get("repeat") or not card.get("dueDate"): return
        nd = next_due(card["dueDate"], card["repeat"])
        if nd is None or (card.get("repeatUntil") and nd.isoformat() > card["repeatUntil"]): card["repeat"] = ""; self.toast("Repeat ended")
        else: self.store.log(card, f"Skipped {fmt_date(card['dueDate'])} → {fmt_date(nd.isoformat())}"); card["dueDate"] = nd.isoformat(); self.toast(f"Skipped — next due {fmt_date(card['dueDate'])}")
        card["updatedAt"] = now_iso(); self.store.save(); self.refresh()

    def manage_labels(self):
        LabelDialog(self, self.store).exec(); self.filter_labels &= {l["id"] for l in self.store.labels}; self.refresh()

    def column_dialog(self, col_id):
        d = ColumnDialog(self, self.store, col_id)
        if d.exec() != QDialog.Accepted: return
        s = self.store; cols = s.columns; a = d.result_action
        if a == "save":
            if d.col: d.col.update(label=d.name.text().strip(), icon=d.icon.text().strip(), color=d.color); self.toast("Column updated")
            else: cols.append({"id": "c_" + uid(), "label": d.name.text().strip(), "icon": d.icon.text().strip(), "color": d.color}); [s.ensure_year(y) for y in s.years()]; self.toast("Column added")
        elif a in ("left", "right"):
            i = cols.index(d.col); j = i + (-1 if a == "left" else 1)
            if 0 <= j < len(cols): cols[i], cols[j] = cols[j], cols[i]
        elif a == "delete":
            n = sum(len(kd.get(d.col["id"], [])) for Y in s.db["years"].values() for kd in Y["kanban"].values())
            if QMessageBox.question(self, "Delete column", f"Delete column \"{d.col['label']}\"?" + (f" Its {n} task(s) across all months will move to To Do." if n else "")) != QMessageBox.Yes: return
            for Y in s.db["years"].values():
                for kd in Y["kanban"].values():
                    arr = kd.pop(d.col["id"], [])
                    for c in arr: c["done"] = False; s.log(c, f"Column \"{d.col['label']}\" deleted — moved to To Do")
                    kd["todo"] = arr + kd.get("todo", [])
            cols.remove(d.col)
        s.save(); self.refresh()

    # ---------- detail panel ----------
    DOCK_W = 440

    def open_detail(self, cid, col_id):
        card, col_id = self.store.find_card(cid, col_id)
        if not card: return
        self.detail_ref = (cid, col_id); self.render_detail()
        if self.dock.isVisible(): return
        self.dock.show()
        if self.motion():       # the board reflows as the pane slides in (setMaximumWidth also lowers the minimum while below it)
            self.dock.setMaximumWidth(1); self.animate(self.dock, b"maximumWidth", 1, self.DOCK_W, 260, on_done=lambda: (self.dock.setMaximumWidth(16777215), self.dock.setMinimumWidth(self.DOCK_W)))

    def close_detail(self):
        self.detail_ref = None
        if self.dock.isVisible() and self.motion():
            self.animate(self.dock, b"maximumWidth", self.dock.width(), 1, 200, QEasingCurve.InCubic, on_done=lambda: (self.dock.hide(), self.dock.setMaximumWidth(16777215), self.dock.setMinimumWidth(self.DOCK_W)))
        else: self.dock.hide()

    def _dcard(self):
        if not self.detail_ref: return None, None
        return self.store.find_card(*self.detail_ref)

    def render_detail(self):
        card, col_id = self._dcard()
        if not card: self.close_detail(); return
        s = self.store; col = s.col(col_id); done = card.get("done") or col_id == "done"
        self._clear(self.detail_layout); L = self.detail_layout; L.setContentsMargins(0, 0, 0, 0)
        inner = QWidget(); v = QVBoxLayout(inner); v.setContentsMargins(20, 16, 20, 16); v.setSpacing(12)
        top = QHBoxLayout(); top.addWidget(badge(f"{col.get('icon', '')} {col['label']}", "#fff", col["color"])); top.addWidget(badge(PRI_LABEL[card.get("priority", "med")], PRI_COLOR[card.get("priority", "med")]))
        if card.get("dueDate"): od = is_overdue(card["dueDate"], done); top.addWidget(badge(("⚠ Overdue " if od else "📅 ") + fmt_date(card["dueDate"]), "#ff6584" if od else "#8890b0"))
        if card.get("repeat"): top.addWidget(badge("↻ " + REPEAT_SHORT.get(card["repeat"], card["repeat"]), "#b48aff"))
        top.addStretch(); cl = QPushButton("✕"); cl.setObjectName("ghost"); cl.clicked.connect(self.close_detail); top.addWidget(cl); v.addLayout(top)
        title = QLineEdit(card["title"]); title.setStyleSheet("font-size:17px;font-weight:700;background:transparent;border-color:transparent;"); title.editingFinished.connect(lambda: self.dset("title", title.text())); v.addWidget(title)
        g = QGridLayout(); colc = QComboBox()
        for c in s.columns: colc.addItem(f"{c.get('icon', '')} {c['label']}", c["id"])
        colc.setCurrentIndex(colc.findData(col_id)); colc.currentIndexChanged.connect(lambda: self.dmove(colc.currentData()))
        pri = QComboBox()
        for k, lab in PRI_LABEL.items(): pri.addItem(lab, k)
        pri.setCurrentIndex(pri.findData(card.get("priority", "med"))); pri.currentIndexChanged.connect(lambda: self.dset("priority", pri.currentData()))
        dcb = QCheckBox("Due"); de = QDateEdit(QDate.currentDate()); de.setCalendarPopup(True); de.setDisplayFormat("yyyy-MM-dd"); de.setMinimumWidth(150)
        if card.get("dueDate"): dcb.setChecked(True); de.setDate(QDate.fromString(card["dueDate"], "yyyy-MM-dd"))
        de.setEnabled(dcb.isChecked()); dcb.toggled.connect(lambda on: self.dset("dueDate", de.date().toString("yyyy-MM-dd") if on else "")); de.dateChanged.connect(lambda d: self.dset("dueDate", d.toString("yyyy-MM-dd")) if dcb.isChecked() else None)
        for i, (lab, w) in enumerate((("Column", colc), ("Priority", pri))): g.addWidget(QLabel(lab, objectName="muted"), 0, i); g.addWidget(w, 1, i)
        dl = QHBoxLayout(); dl.addWidget(dcb); dl.addWidget(de); dl.addStretch(); g.addWidget(QLabel("Due date", objectName="muted"), 2, 0); g.addLayout(dl, 3, 0, 1, 2)
        rp = QComboBox()
        for k, lab in REPEATS: rp.addItem(("↻ " if k else "") + lab, k)
        rp.setCurrentIndex(max(0, rp.findData(card.get("repeat", "")))); rp.currentIndexChanged.connect(lambda: self.dset_repeat(rp.currentData()))
        rl = QHBoxLayout(); rl.addWidget(rp)
        if card.get("repeat"):
            nd = next_due(card["dueDate"], card["repeat"]) if card.get("dueDate") else None
            rl.addWidget(QLabel(("then " + fmt_date(nd.isoformat()) if nd else "") + (f" · until {fmt_date(card['repeatUntil'])}" if card.get("repeatUntil") else ""), objectName="muted"))
            sk = QPushButton("Skip once"); sk.setObjectName("ghost"); sk.setToolTip("Move this task to its next occurrence without completing it"); sk.clicked.connect(lambda: self.skip_occurrence(card["id"], col_id)); rl.addWidget(sk)
        rl.addStretch(); g.addWidget(QLabel("Repeat", objectName="muted"), 4, 0); g.addLayout(rl, 5, 0, 1, 2); v.addLayout(g)
        lh = QHBoxLayout(); lh.addWidget(QLabel("LABELS", objectName="navSection")); lh.addStretch(); mg = QPushButton("Manage"); mg.setObjectName("ghost"); mg.clicked.connect(self.manage_labels); lh.addWidget(mg); v.addLayout(lh)
        lr = QGridLayout(); lr.setSpacing(6)        # ponytail: 3 chips per row instead of a flow layout; wraps within the 440 px dock
        for i, l in enumerate(s.labels):
            b = QPushButton(l["name"]); b.setCheckable(True); b.setChecked(l["id"] in card.get("labels", [])); b.setStyleSheet(f"QPushButton{{border-color:{l['color']};color:{l['color']};padding:3px 10px;}}QPushButton:checked{{background:{l['color']};color:#fff;}}")
            b.toggled.connect(lambda on, lid=l["id"]: self.dlabel(lid, on)); lr.addWidget(b, i // 3, i % 3)
        if not s.labels: lr.addWidget(QLabel("No labels yet", objectName="muted"), 0, 0)
        v.addLayout(lr)
        v.addWidget(QLabel("DESCRIPTION", objectName="navSection")); desc = QTextEdit(card.get("desc", "")); desc.setPlaceholderText("Add details, links, notes…"); desc.setFixedHeight(80)
        desc.focusOutEvent = lambda e, te=desc: (QTextEdit.focusOutEvent(te, e), self.dset("desc", te.toPlainText())); v.addWidget(desc)
        subs = card.get("subs", []); sd = sum(1 for x in subs if x.get("done")); v.addWidget(QLabel(f"CHECKLIST  {sd}/{len(subs)}" if subs else "CHECKLIST", objectName="navSection"))
        for x in subs:
            r = QHBoxLayout(); r.setSpacing(6); cb = QCheckBox(); cb.setChecked(bool(x.get("done"))); cb.toggled.connect(lambda on, sid=x["id"]: self.dsub(sid, on))
            tl = QLabel(x["text"]); tl.setWordWrap(True); tl.setToolTip("Double-click to edit"); tl.setStyleSheet("text-decoration:line-through;color:#5c6382;" if x.get("done") else "")
            tl.mouseDoubleClickEvent = lambda e, sid=x["id"]: self.dsub_edit(sid)
            eb = QPushButton("✎"); eb.setObjectName("ghost"); eb.setToolTip("Edit item"); eb.clicked.connect(lambda _, sid=x["id"]: self.dsub_edit(sid))
            rm = QPushButton("✕"); rm.setObjectName("ghost"); rm.clicked.connect(lambda _, sid=x["id"]: self.dsub_remove(sid))
            r.addWidget(cb, 0, Qt.AlignTop); r.addWidget(tl, 1); r.addWidget(eb, 0, Qt.AlignTop); r.addWidget(rm, 0, Qt.AlignTop); v.addLayout(r)
        ar = QHBoxLayout(); si = QLineEdit(); si.setPlaceholderText("Add an item… (Enter)"); si.returnPressed.connect(lambda: self.dsub_add(si.text())); ab = QPushButton("Add"); ab.clicked.connect(lambda: self.dsub_add(si.text())); ar.addWidget(si); ar.addWidget(ab); v.addLayout(ar)
        comments = card.get("comments", []); v.addWidget(QLabel(f"COMMENTS  {len(comments) or ''}", objectName="navSection"))
        cr = QHBoxLayout(); ci = QTextEdit(); ci.setPlaceholderText("Write a comment… (Ctrl+Enter to post)"); ci.setFixedHeight(56); pb = QPushButton("Post"); pb.setObjectName("primary"); pb.clicked.connect(lambda: self.dcomment(ci.toPlainText()))
        QShortcut(QKeySequence("Ctrl+Return"), ci, activated=lambda: self.dcomment(ci.toPlainText())); cr.addWidget(ci, 1); cr.addWidget(pb, 0, Qt.AlignBottom); v.addLayout(cr)
        for m in reversed(comments):
            f = QFrame(); f.setObjectName("panel"); fl = QVBoxLayout(f); hr = QHBoxLayout(); hr.addWidget(QLabel(fmt_dt(m["at"]), objectName="muted")); hr.addStretch(); dbn = QPushButton("Delete"); dbn.setObjectName("ghost"); dbn.clicked.connect(lambda _, mid=m["id"]: self.dcomment_del(mid)); hr.addWidget(dbn); fl.addLayout(hr)
            t = QLabel(m["text"]); t.setWordWrap(True); fl.addWidget(t); v.addWidget(f)
        v.addWidget(QLabel("ACTIVITY", objectName="navSection"))
        for a in reversed(card.get("activity", [])[-20:]): v.addWidget(QLabel(f"{fmt_dt(a['at'])}   {a['text']}", objectName="muted2", wordWrap=True))
        v.addStretch(); sa = QScrollArea(); sa.setWidgetResizable(True); sa.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff); sa.setWidget(inner); L.addWidget(sa, 1)
        foot = QHBoxLayout(); foot.setContentsMargins(16, 8, 16, 12); mk = QPushButton("↩ Reopen" if done else "✓ Mark complete"); mk.setObjectName("" if done else "primary"); mk.clicked.connect(lambda: self.toggle_done(card["id"], col_id))
        ed = QPushButton("✎ Edit in form"); ed.clicked.connect(lambda: self.edit_task(card["id"], col_id)); dl2 = QPushButton("Delete"); dl2.setObjectName("danger"); dl2.clicked.connect(lambda: self.delete_task(card["id"], col_id))
        foot.addWidget(mk); foot.addWidget(ed); foot.addStretch(); foot.addWidget(dl2); L.addLayout(foot)

    def dset_repeat(self, rep):
        card, _ = self._dcard()
        if not card: return
        if rep and not card.get("dueDate"): card["dueDate"] = dt.date.today().isoformat()
        self.dset("repeat", rep)

    def dset(self, field, val):
        card, col = self._dcard()
        if not card: return
        val = val.strip() if isinstance(val, str) else val
        if field == "title" and not val: self.render_detail(); return
        if (card.get(field) or "") == val: return
        card[field] = val; card["updatedAt"] = now_iso()
        self.store.log(card, {"title": "Title updated", "desc": "Description updated", "priority": f"Priority set to {PRI_LABEL.get(val, val)}", "dueDate": f"Due date set to {fmt_date(val)}" if val else "Due date removed", "repeat": f"Repeats {dict(REPEATS).get(val, val).lower()}" if val else "Repeat removed"}[field])
        self.store.save(); self.refresh()

    def dmove(self, to):
        card, col = self._dcard()
        if not card or to == col: return
        self.store.move_card(card["id"], col, to); self.detail_ref = (card["id"], to); self.store.save(); self.refresh()

    def dlabel(self, lid, on):
        card, col = self._dcard()
        if not card: return
        ls = card.setdefault("labels", []); name = next((l["name"] for l in self.store.labels if l["id"] == lid), "")
        if on and lid not in ls: ls.append(lid); self.store.log(card, f"Label \"{name}\" added")
        elif not on and lid in ls: ls.remove(lid); self.store.log(card, f"Label \"{name}\" removed")
        card["updatedAt"] = now_iso(); self.store.save(); self.refresh()

    def dsub(self, sid, on):
        card, col = self._dcard()
        if not card: return
        for x in card.get("subs", []):
            if x["id"] == sid: x["done"] = on
        if card.get("subs") and all(x.get("done") for x in card["subs"]) and col != "done":
            self.store.move_card(card["id"], col, "done"); self.store.log(card, "All sub-tasks completed"); self.detail_ref = (card["id"], "done"); self.toast("🎉 All sub-tasks done — moved to Completed!")
        self.store.save(); self.refresh()

    def dsub_add(self, t):
        card, col = self._dcard(); t = t.strip()
        if not card or not t: return
        card.setdefault("subs", []).append({"id": uid(), "text": t, "done": False}); self.store.log(card, f"Checklist item added: {t}"); self.store.save(); self.refresh()

    def dsub_edit(self, sid):
        card, col = self._dcard()
        if not card: return
        x = next((x for x in card.get("subs", []) if x["id"] == sid), None)
        if not x: return
        t, ok = QInputDialog.getMultiLineText(self, "Edit checklist item", "Text:", x["text"])
        if ok and t.strip() and t.strip() != x["text"]: x["text"] = t.strip(); card["updatedAt"] = now_iso(); self.store.log(card, "Checklist item edited"); self.store.save(); self.refresh()

    def dsub_remove(self, sid):
        card, col = self._dcard()
        if not card: return
        card["subs"] = [x for x in card.get("subs", []) if x["id"] != sid]; self.store.save(); self.refresh()

    def dcomment(self, t):
        card, col = self._dcard(); t = t.strip()
        if not card or not t: return
        card.setdefault("comments", []).append({"id": uid(), "text": t, "at": now_iso()}); self.store.log(card, "Comment added"); self.store.save(); self.refresh()

    def dcomment_del(self, mid):
        card, col = self._dcard()
        if card: card["comments"] = [m for m in card.get("comments", []) if m["id"] != mid]; self.store.save(); self.refresh()

    # ---------- calendar ----------
    def render_calendar(self):
        s = self.store; y, m = s.year, s.month; self._clear(self.cal_grid); self._clear(self.cal_side)
        first = dt.date(y, m + 1, 1); start = first - dt.timedelta(days=first.weekday()); days_in = calendar.monthrange(y, m + 1)[1]
        cells = -(-(first.weekday() + days_in) // 7) * 7; end = start + dt.timedelta(days=cells - 1); today = dt.date.today().isoformat()
        by_date, unscheduled, other, nghost = {}, [], [], [0]
        def iso_date(v):
            try: return dt.date.fromisoformat(v) if v else None
            except ValueError: return None
        for c in s.columns:
            for k in s.kanban().get(c["id"], []):
                done = k.get("done") or c["id"] == "done"; d = iso_date(k.get("dueDate")); in_grid = d is not None and start <= d <= end
                if not k.get("dueDate") and done: continue
                chip = CalChip(("↻ " if k.get("repeat") else "") + ("" if in_grid or not d else fmt_date(k["dueDate"]) + " · ") + k["title"], c["color"], k["id"], done, is_overdue(k.get("dueDate"), done),
                               on_click=lambda kid=k["id"], cid=c["id"]: self.open_detail(kid, cid))
                (unscheduled if not d else by_date.setdefault(k["dueDate"], []) if in_grid else other).append(chip)
                if not done:
                    for od in occurrences(k, start, end):
                        by_date.setdefault(od.isoformat(), []).append(CalChip("↻ " + k["title"], c["color"], None, dashed=True, ghost=True, on_click=lambda kid=k["id"], cid=c["id"]: self.open_detail(kid, cid))); nghost[0] += 1
        self.cal_stat.setText(f"{sum(len(v) for v in by_date.values()) - nghost[0]} scheduled · {len(unscheduled)} unscheduled" + (f" · {nghost[0]} repeats" if nghost[0] else "") + (f" · {len(other)} due in another month" if other else ""))
        for i, dn in enumerate(("Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun")):
            self.cal_grid.addWidget(QLabel(dn, objectName="muted", alignment=Qt.AlignCenter), 0, i); self.cal_grid.setColumnStretch(i, 1)
        for i in range(cells):
            d = start + dt.timedelta(days=i); iso = d.isoformat(); out = d.month != m + 1
            cell = CalDay(iso); cell.setProperty("out", out); cell.setProperty("today", iso == today); cell.setMinimumHeight(96)
            cell.dropped.connect(self.cal_drop); cell.added.connect(lambda iso: self.add_task("todo", dict(dueDate=iso)))
            cv = QVBoxLayout(cell); cv.setContentsMargins(6, 4, 6, 6); cv.setSpacing(3); cv.setAlignment(Qt.AlignTop); hr = QHBoxLayout()
            num = QLabel(str(d.day) + (" " + MONTHS[d.month - 1] if d.day == 1 else "")); num.setStyleSheet("font-weight:700;" + ("color:#5c6382;" if out else "")); hr.addWidget(num); hr.addStretch()
            ab = QPushButton("+"); ab.setObjectName("ghost"); ab.setFixedSize(22, 22); ab.setToolTip(f"Add a task due {fmt_date(iso)}"); ab.clicked.connect(lambda _, iso=iso: self.add_task("todo", dict(dueDate=iso))); hr.addWidget(ab); cv.addLayout(hr)
            for ch in by_date.get(iso, []): cv.addWidget(ch)
            self.cal_grid.addWidget(cell, 1 + i // 7, i % 7)
        tray = CalDay(""); tray.dropped.connect(self.cal_drop); tv = QVBoxLayout(tray); tv.setAlignment(Qt.AlignTop); tv.addWidget(QLabel(f"UNSCHEDULED   {len(unscheduled)}", objectName="navSection"))
        for ch in unscheduled: tv.addWidget(ch)
        if not unscheduled: tv.addWidget(QLabel("Every open task has a date. Drop a task here to clear its due date.", objectName="muted", wordWrap=True))
        self.cal_side.addWidget(tray)
        if other:
            f = QFrame(); f.setObjectName("panel"); fv = QVBoxLayout(f); fv.addWidget(QLabel(f"DUE IN ANOTHER MONTH   {len(other)}", objectName="navSection"))
            for ch in other: fv.addWidget(ch)
            self.cal_side.addWidget(f)

    def cal_drop(self, cid, iso):
        card, _ = self.store.find_card(cid)
        if not card or (card.get("dueDate") or "") == iso: return
        card["dueDate"] = iso; card["updatedAt"] = now_iso(); self.store.log(card, f"Rescheduled to {fmt_date(iso)}" if iso else "Due date removed"); self.store.save(); self.refresh()
        self.toast(f"“{card['title']}” now due {fmt_date(iso)}" if iso else f"Due date cleared for “{card['title']}”")

    # ---------- list (same cards as the board, one line each) ----------
    def render_list(self, *_):
        s = self.store; self._clear(self.list_layout); kd = s.kanban(); q = self.list_search.text().strip().lower(); hide = self.list_hide_done.isChecked(); mode = self.list_group.currentData(); today = dt.date.today()
        rows = [(c, col) for col in s.columns for c in kd.get(col["id"], []) if not (hide and (c.get("done") or col["id"] == "done"))
                and (not q or q in " ".join([c.get("title", ""), c.get("desc", "")] + [x["text"] for x in c.get("subs", [])]).lower())]
        def key(c, col):
            if mode == "column": return (s.columns.index(col), f"{col.get('icon', '')} {col['label']}", col["color"])
            if mode == "priority": p = c.get("priority", "med"); return ({"high": 0, "med": 1, "low": 2}[p], PRI_LABEL[p], PRI_COLOR[p])
            if mode == "label":
                l = next((x for x in s.labels if x["id"] in c.get("labels", [])), None)
                return (s.labels.index(l), l["name"], l["color"]) if l else (999, "No label", "#8890b0")
            try: n = (dt.date.fromisoformat(c["dueDate"]) - today).days
            except (KeyError, TypeError, ValueError): return (6, "No due date", "#8890b0")
            if n < 0: return (0, "Overdue", "#ff6584") if not (c.get("done") or col["id"] == "done") else (5, "Past", "#8890b0")
            return (1, "Today", "#ffb347") if n == 0 else (2, "Tomorrow", "#6c8aff") if n == 1 else (3, "This week", "#6c8aff") if n <= 6 else (4, "Next 30 days", "#8890b0") if n <= 30 else (5, "Later", "#8890b0")
        groups = {}
        for c, col in rows: groups.setdefault(key(c, col), []).append((c, col))
        for (order, name, color) in sorted(groups):
            items = groups[(order, name, color)]; hd = QHBoxLayout(); t = QLabel(name); t.setStyleSheet(f"font-weight:700;color:{color};margin-top:8px;"); hd.addWidget(t); hd.addWidget(badge(str(len(items)), color)); hd.addStretch(); self.list_layout.addLayout(hd)
            for c, col in items:
                r = ListRow(s, c, col); r.clicked.connect(self.open_detail); r.toggle.connect(self.toggle_done); self.list_layout.addWidget(r)
        if not rows: self.list_layout.addWidget(QLabel("No tasks match." if q or hide else "No tasks this month. Add one with + Add Task or Ctrl N.", objectName="muted"))
        ndone = sum(1 for c, col in rows if c.get("done") or col["id"] == "done"); self.list_stat.setText(f"{len(rows)} tasks · {ndone} done")

    # ---------- rollover ----------
    def open_rollover(self):
        s = self.store; d = QDialog(self); d.setWindowTitle(f"Month Rollover · {s.year}"); d.setMinimumWidth(560); v = QVBoxLayout(d)
        v.addWidget(QLabel("Auto-rollover is ON: incomplete tasks from past months move forward each time the app opens. Completed tasks stay where they were finished. December → January of the next year.", wordWrap=True, objectName="muted2"))
        for mi in range(12):
            inc = s.incomplete_cards(mi, s.year); tot = len(inc); past = s.month_past(mi, s.year)
            ty, tm = (s.year + 1, 0) if mi == 11 else (s.year, mi + 1)
            r = QHBoxLayout(); r.addWidget(QLabel(f"{MONTHS[mi]} {s.year} → {MONTHS[tm]} {ty}", objectName="sectitle"))
            st = "🔒 Not ended yet" if not past else ("✅ All done" if tot == 0 else f"⏳ {tot} pending"); r.addWidget(QLabel(st, objectName="muted")); r.addStretch()
            if past and tot:
                b = QPushButton(f"Move → {MONTHS[tm]} {ty}"); b.setObjectName("primary")
                b.clicked.connect(lambda _, m=mi, dd=d: (s.migrate_month(m, s.year), self.toast("✅ Moved"), dd.accept(), self.refresh(), self.open_rollover())); r.addWidget(b)
            v.addLayout(r)
        hist = s.db["migrations"][-8:][::-1]
        if hist:
            v.addWidget(QLabel("HISTORY", objectName="navSection"))
            for m in hist: v.addWidget(QLabel(f"{'🤖' if m.get('auto') else '👆'} {MONTHS[m['from']]} {m.get('fromYear')} → {MONTHS[m['to']]} {m.get('toYear')}   +{m['kanban'] + m.get('checklist', 0)} tasks   {fmt_dt(m['at'])}", objectName="muted2"))
        cb = QPushButton("Close"); cb.clicked.connect(d.accept); v.addWidget(cb); d.exec()

    # ---------- backup ----------
    def backup_now_ui(self):
        self.store.save(); dest = self.store.backup_now()
        if dest: self.toast("💾 Backup saved: " + os.path.basename(dest))

    def open_backup(self):
        d = QDialog(self); d.setWindowTitle("Backup & Restore"); d.setMinimumWidth(600); v = QVBoxLayout(d)
        v.addWidget(QLabel("Automatic every 30 min and on quit · last 30 kept in the backup folder", objectName="muted"))
        v.addWidget(QLabel("BACKUP FOLDER", objectName="navSection")); cur = get_backup_dir(); pl = QLabel(cur); pl.setWordWrap(True); pl.setObjectName("muted2"); v.addWidget(pl)
        r = QHBoxLayout()
        def choose():
            p = QFileDialog.getExistingDirectory(d, "Choose backup folder", cur)
            if p: st = read_settings(); st["backupDir"] = p; write_settings(st); d.accept(); self.toast("Backup folder: " + p); self.open_backup()
        def set_dir(p):
            st = read_settings()
            if p: st["backupDir"] = p
            else: st.pop("backupDir", None)
            write_settings(st); d.accept(); self.toast("Backups now go to " + get_backup_dir()); self.open_backup()
        cb = QPushButton("📁 Choose folder…"); cb.setObjectName("primary"); cb.clicked.connect(choose); r.addWidget(cb)
        od = detect_onedrive()
        if od: ob = QPushButton("☁ Use OneDrive"); ob.setToolTip(os.path.join(od, f"{APP_NAME} Backups")); ob.clicked.connect(lambda: set_dir(os.path.join(od, f"{APP_NAME} Backups"))); r.addWidget(ob)
        else: r.addWidget(QLabel("OneDrive not detected — use Choose folder…", objectName="muted"))
        if cur != DEFAULT_BACKUP_DIR: rb = QPushButton("Reset to default"); rb.clicked.connect(lambda: set_dir(None)); r.addWidget(rb)
        ofb = QPushButton("Open folder"); ofb.clicked.connect(lambda: self.open_path(cur)); r.addWidget(ofb); r.addStretch(); v.addLayout(r)
        lst = self.store.backup_list(); v.addWidget(QLabel(f"BACKUPS IN THIS FOLDER ({len(lst)})", objectName="navSection"))
        r2 = QHBoxLayout(); bn = QPushButton("💾 Backup now"); bn.setObjectName("primary"); bn.clicked.connect(lambda: (self.backup_now_ui(), d.accept(), self.open_backup())); r2.addWidget(bn)
        def restore_file():
            p, _ = QFileDialog.getOpenFileName(d, "Restore from backup file", cur, "Backup JSON (*.json)")
            if p: self._restore(p, d)
        rf = QPushButton("↩ Restore from file…"); rf.clicked.connect(restore_file); r2.addWidget(rf); r2.addStretch(); v.addLayout(r2)
        lw = QListWidget(); lw.setFixedHeight(200)
        for name, size, mt in lst: lw.addItem(f"{name}    {dt.datetime.fromtimestamp(mt).strftime('%Y-%m-%d %H:%M')}    {size / 1024:.1f} KB")
        v.addWidget(lw); rs = QPushButton("Restore selected"); rs.clicked.connect(lambda: self._restore(os.path.join(cur, lst[lw.currentRow()][0]), d) if lw.currentRow() >= 0 else None); v.addWidget(rs)
        v.addWidget(QLabel("LIVE DATA FILE", objectName="navSection")); v.addWidget(QLabel(DB_FILE, objectName="muted2", wordWrap=True))
        cl = QPushButton("Close"); cl.clicked.connect(d.accept); v.addWidget(cl); d.exec()

    def _restore(self, path, dlg=None):
        if QMessageBox.warning(self, "Restore backup", f"Restore from:\n{path}\n\nThis will overwrite current data (a safety backup is taken first). Continue?", QMessageBox.Yes | QMessageBox.Cancel) != QMessageBox.Yes: return
        try:
            self.store.restore_file(path)
        except Exception as e:
            QMessageBox.critical(self, "Restore failed", str(e)); return
        if dlg: dlg.accept()
        self.close_detail(); self.refresh(); self.toast("✅ Restored from " + os.path.basename(path))

    # ---------- export ----------
    def export_excel(self):
        d = ExportDialog(self, self.store)
        if d.exec() != QDialog.Accepted or not d.months(): return
        tag = MONTHS[d.months()[0]] if len(d.months()) == 1 else f"{len(d.months())}months"
        p, _ = QFileDialog.getSaveFileName(self, "Export Excel", os.path.join(os.path.expanduser("~"), "Desktop", f"{APP_NAME}_{tag}_{self.store.year}.xlsx"), "Excel (*.xlsx)")
        if not p: return
        try:
            export_excel(self.store, p, d.months(), d.inc_board.isChecked(), d.inc_sum.isChecked(), d.inc_done.isChecked())
        except Exception as e:
            QMessageBox.critical(self, "Export failed", str(e)); return
        self.toast("📊 Exported: " + os.path.basename(p)); self.open_path(os.path.dirname(p))

    # ---------- appearance ----------
    def palette_dict(self):
        a = self.appearance; p = dict(THEMES[a.get("theme", "dark")])
        if a.get("bg"): p["bg"] = a["bg"]
        if a.get("surface"): p["s1"] = p["s2"] = a["surface"]
        if a.get("text"): p["text"] = a["text"]
        if a.get("accent"): p["accent"] = a["accent"]
        return p

    def apply_appearance(self):
        a = self.appearance; p = self.palette_dict()
        fam = a.get("font") or ("Segoe UI" if sys.platform.startswith("win") else QApplication.font().family())
        QApplication.instance().setStyleSheet(build_qss(p, fam, int(a.get("size", 13))))
        self.theme_btn.setText("☀️" if a.get("theme") == "light" else "🌙"); self.chart.accent = p["accent"]; self.chart.update()

    def save_appearance(self):
        st = read_settings(); st["appearance"] = self.appearance; write_settings(st)
        self.store.db["meta"]["appearance_py"] = self.appearance; self.store.save(); self.apply_appearance()

    def toggle_theme(self):
        self.appearance["theme"] = "light" if self.appearance.get("theme") == "dark" else "dark"; self.save_appearance()

    def open_appearance(self):
        d = QDialog(self); d.setWindowTitle("Appearance"); d.setMinimumWidth(520); v = QVBoxLayout(d); a = self.appearance
        v.addWidget(QLabel("FONT", objectName="navSection")); fr = QHBoxLayout(); fc = QFontComboBox()
        if a.get("font"): fc.setCurrentFont(QFont(a["font"]))
        fc.currentFontChanged.connect(lambda f: (a.__setitem__("font", f.family()), self.save_appearance())); fr.addWidget(QLabel("Family")); fr.addWidget(fc, 1)
        sz = QSpinBox(); sz.setRange(10, 20); sz.setValue(int(a.get("size", 13))); sz.valueChanged.connect(lambda n: (a.__setitem__("size", n), self.save_appearance())); fr.addWidget(QLabel("Size")); fr.addWidget(sz); v.addLayout(fr)
        v.addWidget(QLabel("COLOURS", objectName="navSection"))
        def crow(label, key, varname):
            r = QHBoxLayout(); r.addWidget(QLabel(label), 1); cur = a.get(key) or self.palette_dict()[varname]; b = QPushButton(cur); b.setStyleSheet(f"background:{cur};color:#fff;")
            def pick():
                c = QColorDialog.getColor(QColor(cur), d, label)
                if c.isValid(): a[key] = c.name(); self.save_appearance(); d.accept(); self.open_appearance()
            b.clicked.connect(pick); r.addWidget(b)
            if a.get(key): rb = QPushButton("Reset"); rb.setObjectName("ghost"); rb.clicked.connect(lambda: (a.__setitem__(key, ""), self.save_appearance(), d.accept(), self.open_appearance())); r.addWidget(rb)
            v.addLayout(r)
        crow("Text colour", "text", "text"); crow("Accent colour (buttons, highlights)", "accent", "accent"); crow("Background", "bg", "bg"); crow("Cards & panels", "surface", "s2")
        tr = QHBoxLayout(); tr.addWidget(QLabel("Dark / light base theme"), 1); tb = QPushButton(("☀️ Light" if a.get("theme") == "light" else "🌙 Dark") + " — switch"); tb.clicked.connect(lambda: (self.toggle_theme(), d.accept(), self.open_appearance())); tr.addWidget(tb); v.addLayout(tr)
        mr = QHBoxLayout(); mr.addWidget(QLabel("Animations & transitions"), 1); mc = QCheckBox("On"); mc.setChecked(self.motion()); mc.toggled.connect(lambda on: (a.__setitem__("motion", on), self.save_appearance())); mr.addWidget(mc); v.addLayout(mr)
        v.addWidget(QLabel("PRESETS", objectName="navSection")); pg = QGridLayout()
        for i, (name, theme, vals) in enumerate(PRESETS):
            b = QPushButton(name); b.setStyleSheet(f"background:{vals.get('bg', THEMES[theme]['bg'])};color:{vals.get('text', THEMES[theme]['text'])};border-color:{vals.get('accent', '#6c8aff')};")
            b.clicked.connect(lambda _, t=theme, vv=vals: (a.update(theme=t, bg=vv.get("bg", ""), surface=vv.get("surface", ""), text=vv.get("text", ""), accent=vv.get("accent", "")), self.save_appearance(), d.accept(), self.open_appearance())); pg.addWidget(b, i // 4, i % 4)
        v.addLayout(pg)
        rr = QHBoxLayout(); ra = QPushButton("Reset everything to default"); ra.clicked.connect(lambda: (a.update(theme="dark", font="", size=13, text="", accent="", bg="", surface="", motion=True), self.save_appearance(), d.accept(), self.open_appearance())); rr.addWidget(ra); rr.addStretch(); cl = QPushButton("Close"); cl.clicked.connect(d.accept); rr.addWidget(cl); v.addLayout(rr)
        d.exec()


def main():
    app = QApplication(sys.argv); app.setApplicationName(APP_NAME); app.setQuitOnLastWindowClosed(False)
    store = Store(); w = MainWindow(store); w.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()

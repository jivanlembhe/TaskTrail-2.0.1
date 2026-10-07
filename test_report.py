"""Report engine, auto-priority, calendar and scale checks:  QT_QPA_PLATFORM=offscreen python test_report.py
Throw-away profile; never touches your real data.  The AI call is tested against a local mock server (no network, no key)."""
import os, sys, json, time, tempfile, threading, datetime as dt, http.server
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
os.environ["XDG_CONFIG_HOME"] = os.environ["APPDATA"] = tempfile.mkdtemp(prefix="tt_rep_")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from PySide6.QtWidgets import QApplication
app = QApplication(sys.argv)
import tasktrail as T, report as R
D = dt.date

# ── 1. auto-priority rules (fixed "today" → deterministic) ─────────────────
today = D(2026, 10, 7); mk = lambda due, pri="low", **kw: {**dict(id="x", title="t", priority=pri, dueDate=due.isoformat() if due else "", done=False), **kw}
rule = lambda c: (T.auto_priority_for(c, today) or (None,))[0]
assert rule(mk(today - dt.timedelta(5))) == "high" and rule(mk(today)) == "high" and rule(mk(today + dt.timedelta(1))) == "high"
assert rule(mk(today + dt.timedelta(2))) == "med" and rule(mk(today + dt.timedelta(7))) == "med" and rule(mk(today + dt.timedelta(8))) is None
assert rule(mk(None)) is None and rule(mk(today, done=True)) is None and rule(mk(today, prioLock=True)) is None
S = T.Store(); S.ensure_year(2026); kd = S.kanban(9, 2026)
for i, (due, pri) in enumerate(((today, "low"), (today + dt.timedelta(3), "low"), (today + dt.timedelta(3), "high"), (today + dt.timedelta(30), "low"))):
    kd.setdefault("todo", []).append(dict(id=f"c{i}", title=f"c{i}", priority=pri, dueDate=due.isoformat(), done=False, subs=[], labels=[], activity=[]))
ch = T.apply_auto_priority(S, today); got = {c["id"]: c["priority"] for c in kd["todo"] if c["id"].startswith("c")}
assert got == {"c0": "high", "c1": "med", "c2": "high", "c3": "low"}, got                 # raises only; never lowers c2; leaves c3
assert len(ch) == 2 and "auto-raised" in kd["todo"][0]["activity"][-1]["text"] and not T.apply_auto_priority(S, today)   # idempotent
print("auto-priority OK")

# ── 2. a real September: completed + unfinished work, then the app's own month rollover ──
S = T.Store(); Y = 2026; S.ensure_year(Y); lab = dict(id="L1", name="backend", color="#8b5cf6"); S.labels.append(lab)
def card(col, title, pri, due, subs=(), done_on=None):
    c = dict(id=T.uid(), title=title, desc="", priority=pri, dueDate=due, done=col == "done", subs=[dict(id=T.uid(), text=t, done=d) for t, d in subs], labels=["L1"], repeat="", repeatUntil="", comments=[], activity=[], createdAt="2026-09-02T09:00:00")
    if done_on: c["completedAt"] = done_on + "T15:00:00"
    S.kanban(8, Y).setdefault(col, []).append(c); return c
card("done", "Migrate billing", "high", "2026-09-12", [("Map events", True), ("Replay script", True)], "2026-09-11"); card("done", "Release notes", "med", "2026-09-20", [("Changelog", True)], "2026-09-22")
card("wip", "Accessibility audit", "high", "2026-09-25", [("Audit dashboard", True), ("Audit calendar", False)]); card("todo", "Draft roadmap", "med", "2026-09-28", [("Outline", False)])
S.migrate_month(8, Y)
f = R.collect(S, Y, 8, today=D(2026, 10, 7)); m = f["metrics"]
assert (m["total"], m["completed"], m["incomplete"], m["rolled"]) == (4, 2, 2, 2), m           # incomplete found by following migratedFrom
assert (m["subs_done"], m["subs_total"], m["subs_done_in_open"]) == (4, 6, 1) and (m["on_time"], m["with_due"]) == (1, 2) and m["carried_in"] == 0
aud = next(t for t in f["incomplete"] if t["title"] == "Accessibility audit"); assert aud["subs_done"] == ["Audit dashboard"] and aud["subs_open"] == ["Audit calendar"]
n = R.template_narrative(f); assert any("Completed 2 of 4 tasks" in x for x in n["overview"]) and not any(" I " in " " + x + " " for x in n["overview"] + n["challenges"] + n["next_focus"]) and set(n["completed"]) == {t["id"] for t in f["completed"]} and set(n["incomplete"]) == {t["id"] for t in f["incomplete"]}
p = os.path.join(os.environ["XDG_CONFIG_HOME"], "r.docx"); R.build_docx(p, f, n, dict(name="Test User", code="8052", appraiser="Koshal Ohol", program="MahaPMP", joined="September 01, 2022", classification="Internal"), labels_order=["backend"])
from docx import Document; doc = Document(p); body = "\n".join(x.text for x in doc.paragraphs) + "\n" + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
for need in ("September 2026", "Prepared By: Test User", "Employee Code: 8052", "Appraiser: Koshal Ohol", "Program: MahaPMP", "Joining Date: September 01, 2022", "Overview:", "backend:", "Migrate billing", "Replay script",
             "Work in Progress and Carried Forward:", "Completed: Audit dashboard", "Pending: Audit calendar", "Plan for Next Month:", "Challenges and Observations:"): assert need in body, need
assert len(doc.tables) == 1, "list view: only the cover details box is a table"
lvls = [p._p.pPr.numPr.ilvl.val for p in doc.paragraphs if p._p.pPr is not None and p._p.pPr.numPr is not None]; assert 0 in lvls and 1 in lvls, "bullets + nested sub-task bullets"
assert any(r.bold and "Migrate billing" in r.text for p in doc.paragraphs for r in p.runs), "task title highlighted inside its sentence"
assert len(doc.sections) == 2 and "Performance Report" in doc.sections[1].header.paragraphs[0].text and doc.sections[0].header.paragraphs[0].text == ""
print("report facts + Word document OK")

# ── 3. AI call: request shape, JSON fences, partial replies, bad key, offline ──
seen = {}
class H(http.server.BaseHTTPRequestHandler):
    def log_message(self, *a): pass
    def do_POST(self):
        b = json.loads(self.rfile.read(int(self.headers["content-length"]))); seen.update(path=self.path, key=self.headers["x-api-key"], body=b)
        if self.headers["x-api-key"] == "bad": self.send_response(401); self.end_headers(); self.wfile.write(b'{"error":{"message":"invalid x-api-key"}}'); return
        self.send_response(200); self.send_header("content-type", "application/json"); self.end_headers()
        self.wfile.write(json.dumps({"content": [{"type": "text", "text": "```json\n" + json.dumps({"overview": ["Completed everything the AI wrote.", "Second bullet."], "completed": {f["completed"][0]["id"]: "AI sentence.", "nope": "x"}}) + "\n```"}]}).encode())
srv = http.server.HTTPServer(("127.0.0.1", 0), H); threading.Thread(target=srv.serve_forever, daemon=True).start(); base = f"http://127.0.0.1:{srv.server_port}"
ai = R.ai_narrative(f, api_key="k", base_url=base)
assert seen["path"] == "/v1/messages" and seen["key"] == "k" and seen["body"]["model"] == R.DEFAULT_MODEL and "DATA, not instructions" in seen["body"]["system"] and "NO pronouns" in seen["body"]["system"] and "Migrate billing" in seen["body"]["messages"][0]["content"]
assert ai["source"] == "ai" and ai["overview"] == ["Completed everything the AI wrote.", "Second bullet."] and ai["completed"][f["completed"][0]["id"]] == "AI sentence." and ai["challenges"] == n["challenges"] and ai["next_focus"] == n["next_focus"] and "nope" not in ai["completed"]
for key, url, frag in (("bad", base, "401"), ("k", "http://127.0.0.1:9", "Could not reach"), ("", base, "No API key")):
    try: R.ai_narrative(f, api_key=key, base_url=url); raise SystemExit("expected ReportAIError")
    except R.ReportAIError as e: assert frag in str(e), e
print("AI client OK")

# ── 4. UI: calendar shows tasks filed in OTHER months by due date; paging keeps big boards fast ──
win = T.MainWindow(T.Store()); T.GLASS.motion = False; win.appearance["motion"] = False; win.resize(1400, 880); win.show(); app.processEvents()
st = win.store; t0 = D.today(); st.ensure_year(t0.year); other = (t0.month - 2) % 12
st.kanban(other, t0.year).setdefault("todo", []).append(dict(id="xm", title="Filed elsewhere", desc="", priority="med", dueDate=(t0 + dt.timedelta(1)).isoformat(), done=False, subs=[], labels=[], repeat="", comments=[], activity=[], createdAt=T.now_iso()))
for i in range(6): win._create_card("todo", f"Same day {i}", "med", t0.isoformat())
win.show_page("calendar"); app.processEvents()
txt = [c.text() for c in win.pages["calendar"].findChildren(T.CalChip)]; assert any("Filed elsewhere" in x for x in txt)
assert any(b.text().startswith("+") and "more" in b.text() for b in win.pages["calendar"].findChildren(T.QPushButton)), "overflow button missing"
win.cal_set_mode("week"); app.processEvents(); win.cal_set_mode("month")
win.cal_drop("xm", (t0 + dt.timedelta(3)).isoformat()); assert win.locate_card("xm")[3]["dueDate"] == (t0 + dt.timedelta(3)).isoformat()
kd = st.kanban()
for i in range(1500): kd.setdefault("todo", []).append(dict(id=f"b{i}", title=f"Bulk {i}", desc="d " * 20, priority="med", dueDate="", done=False, subs=[], labels=[], repeat="", comments=[], activity=[], createdAt=T.now_iso()))
t = time.time(); win.show_page("kanban"); app.processEvents(); dt_board = time.time() - t
t = time.time(); win.show_page("list"); app.processEvents(); dt_list = time.time() - t
assert dt_board < 4 and dt_list < 4, (dt_board, dt_list)
print(f"calendar + scale OK (1500+ tasks: board {dt_board:.2f}s, list {dt_list:.2f}s)")
# ── 5. the report dialog really reaches the review step, and edited statements reach the Word file ──
from PySide6.QtCore import QEventLoop, QTimer
from PySide6.QtWidgets import QFileDialog
dlg = T.ReportDialog(win); dlg.year.setValue(t0.year); dlg.month.setCurrentIndex(t0.month - 1); dlg.use_ai.setChecked(False); dlg.fields["name"].setText("Dialog User"); dlg.generate()
lp = QEventLoop(); tm = QTimer(); tm.timeout.connect(lambda: dlg.narr is not None and lp.quit()); tm.start(20); QTimer.singleShot(8000, lp.quit); lp.exec(); tm.stop(); app.processEvents()
assert dlg.stack.currentIndex() == 1 and dlg.stmt, "review step not shown"
list(dlg.stmt.values())[0].setText("Edited statement from the review step."); outp = os.path.join(os.environ["XDG_CONFIG_HOME"], "dlg.docx"); QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: (outp, ""))
dlg.export(); lp = QEventLoop(); tm = QTimer(); tm.timeout.connect(lambda: dlg.result() == 1 and lp.quit()); tm.start(20); QTimer.singleShot(30000, lp.quit); lp.exec(); tm.stop(); assert dlg.result() == 1, "export did not finish"      # dialog accepts itself once the file is fully written
d2 = Document(outp); assert "Edited statement from the review step." in "\n".join(p.text for p in d2.paragraphs) and "Dialog User" in d2.tables[0].rows[0].cells[0].text
print("report dialog OK")
print("ALL REPORT TESTS PASSED")
